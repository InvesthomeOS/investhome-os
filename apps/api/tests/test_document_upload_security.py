"""Document upload validation: magic bytes, MIME/extension match, size, filenames, malware hook."""

from __future__ import annotations

import io

import pytest
from docx import Document as DocxDocument
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image

from investhome_api.config.settings import Settings, get_settings
from investhome_api.services.document_validation import (
    generate_storage_key,
    sanitize_filename,
    validate_upload_content,
)
from investhome_api.services.malware_scan import (
    FailClosedMalwareScanner,
    StubMalwareScanner,
    get_malware_scanner,
    scan_upload_or_raise,
)
from investhome_api.services.storage.factory import get_storage_provider

pytestmark = pytest.mark.usefixtures("client")

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_TEST_DB = "postgresql+psycopg://unit-test-db-user:unit-test-db-password-not-used@localhost:5432/unit_test"
_TEST_JWT = "unit-test-jwt-secret-not-for-production-32"


def _pdf() -> bytes:
    return b"%PDF-1.4\n1 0 obj<< /Type /Catalog >>endobj\ntrailer<<>>\n%%EOF\n"


def _png() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), (40, 80, 120)).save(buf, format="PNG")
    return buf.getvalue()


def _jpeg() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (8, 8), (200, 40, 40)).save(buf, format="JPEG")
    return buf.getvalue()


def _docx() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("Investhome OS supported Word document.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx() -> bytes:
    wb = Workbook()
    wb.active["A1"] = "Amount"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _mz_executable() -> bytes:
    return b"MZ" + b"\x00" * 64 + b"This program cannot be run in DOS mode."


def _post_upload(client: TestClient, filename: str, content: bytes, mime: str) -> dict:
    response = client.post(
        "/documents/upload",
        files={"files": (filename, io.BytesIO(content), mime)},
    )
    assert response.status_code == 201
    return response.json()["results"][0]


def test_valid_pdf_accepted(client: TestClient) -> None:
    result = _post_upload(client, "contract.pdf", _pdf(), "application/pdf")
    assert result["success"] is True
    assert result["document"]["file_extension"] == "pdf"


def test_valid_jpeg_accepted(client: TestClient) -> None:
    result = _post_upload(client, "photo.jpg", _jpeg(), "image/jpeg")
    assert result["success"] is True
    assert result["document"]["file_extension"] == "jpg"


def test_valid_png_accepted(client: TestClient) -> None:
    result = _post_upload(client, "scan.png", _png(), "image/png")
    assert result["success"] is True
    assert result["document"]["file_extension"] == "png"


def test_supported_office_docx_accepted(client: TestClient) -> None:
    result = _post_upload(client, "brief.docx", _docx(), _DOCX_MIME)
    assert result["success"] is True
    assert result["document"]["file_extension"] == "docx"


def test_supported_office_xlsx_accepted(client: TestClient) -> None:
    result = _post_upload(client, "budget.xlsx", _xlsx(), _XLSX_MIME)
    assert result["success"] is True
    assert result["document"]["file_extension"] == "xlsx"


def test_fake_pdf_rejected(client: TestClient) -> None:
    result = _post_upload(client, "spoof.pdf", b"This is not a PDF file.", "application/pdf")
    assert result["success"] is False
    assert result["error"] in {
        "documents.errors.content_mismatch",
        "documents.errors.unknown_type",
    }


def test_executable_renamed_pdf_rejected(client: TestClient) -> None:
    result = _post_upload(client, "payload.pdf", _mz_executable(), "application/pdf")
    assert result["success"] is False
    assert result["error"] == "documents.errors.blocked_type"


def test_mismatched_mime_rejected(client: TestClient) -> None:
    result = _post_upload(client, "photo.jpg", _jpeg(), "application/pdf")
    assert result["success"] is False
    assert result["error"] == "documents.errors.mime_mismatch"


def test_mismatched_extension_vs_content_rejected(client: TestClient) -> None:
    result = _post_upload(client, "report.pdf", _jpeg(), "application/pdf")
    assert result["success"] is False
    assert result["error"] == "documents.errors.content_mismatch"


def test_unknown_binary_rejected(client: TestClient) -> None:
    result = _post_upload(client, "blob.pdf", bytes(range(256)), "application/pdf")
    assert result["success"] is False
    assert result["error"] in {
        "documents.errors.unknown_type",
        "documents.errors.content_mismatch",
        "documents.errors.blocked_type",
    }


def test_traversal_filename_sanitized(client: TestClient) -> None:
    result = _post_upload(client, "../../etc/passwd.pdf", _pdf(), "application/pdf")
    assert result["success"] is True
    stored_name = result["document"]["original_file_name"]
    assert ".." not in stored_name
    assert "/" not in stored_name
    assert "\\" not in stored_name
    assert stored_name.endswith(".pdf")
    stored_key, _ = generate_storage_key("pdf")
    assert ".." not in stored_key
    assert stored_key != "../../etc/passwd.pdf"


def test_oversized_upload_rejected(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("DOCUMENT_MAX_UPLOAD_BYTES", "100")
    get_settings.cache_clear()
    get_storage_provider.cache_clear()
    result = _post_upload(client, "large.txt", b"x" * 200, "text/plain")
    assert result["success"] is False
    assert result["error"] == "documents.errors.file_too_large"


def test_sanitize_filename_strips_traversal() -> None:
    assert sanitize_filename("../../etc/passwd.pdf") == "passwd.pdf"
    assert sanitize_filename(r"..\..\windows\system32\evil.pdf") == "evil.pdf"
    assert ".." not in sanitize_filename("....//....//ok.pdf")
    assert sanitize_filename("..") == "upload"
    assert sanitize_filename("") == "upload"


def test_octet_stream_is_not_trusted() -> None:
    with pytest.raises(HTTPException) as exc:
        validate_upload_content(
            b"MZ" + b"\x00" * 32,
            filename="notes.txt",
            declared_mime="application/octet-stream",
        )
    assert exc.value.detail == "documents.errors.blocked_type"


def test_html_rejected_even_as_txt() -> None:
    with pytest.raises(HTTPException) as exc:
        validate_upload_content(
            b"<html><script>alert(1)</script></html>",
            filename="page.txt",
            declared_mime="text/plain",
        )
    assert exc.value.detail == "documents.errors.blocked_type"


def _settings(**overrides: object) -> Settings:
    kwargs: dict[str, object] = {
        "API_ENVIRONMENT": "development",
        "JWT_SECRET": _TEST_JWT,
        "DATABASE_URL": _TEST_DB,
        "DOCUMENT_MALWARE_SCAN_ENABLED": False,
        "DOCUMENT_MALWARE_SCAN_PROVIDER": "stub",
        "DOCUMENT_MALWARE_SCAN_FAIL_CLOSED": True,
    }
    kwargs.update(overrides)
    return Settings(_env_file=None, **kwargs)


def test_local_malware_stub_allows_upload(monkeypatch: pytest.MonkeyPatch) -> None:
    settings = _settings(DOCUMENT_MALWARE_SCAN_ENABLED=True, API_ENVIRONMENT="development")
    monkeypatch.setattr("investhome_api.services.malware_scan.get_settings", lambda: settings)
    scanner = get_malware_scanner()
    assert isinstance(scanner, StubMalwareScanner)
    verdict = scan_upload_or_raise(b"hello text", filename="note.txt")
    assert verdict.clean is True
    assert verdict.provider == "stub"


def test_production_fail_closed_when_malware_scan_enabled(monkeypatch: pytest.MonkeyPatch) -> None:
    settings = _settings(
        API_ENVIRONMENT="production",
        AUTH_COOKIE_SECURE=True,
        API_DEBUG=False,
        API_ENABLE_OPENAPI=False,
        API_AUTH_ENABLED=True,
        DOCUMENT_MALWARE_SCAN_ENABLED=True,
        DOCUMENT_MALWARE_SCAN_PROVIDER="stub",
        DOCUMENT_MALWARE_SCAN_FAIL_CLOSED=True,
    )
    monkeypatch.setattr("investhome_api.services.malware_scan.get_settings", lambda: settings)
    scanner = get_malware_scanner()
    assert isinstance(scanner, FailClosedMalwareScanner)
    with pytest.raises(HTTPException) as exc:
        scan_upload_or_raise(b"hello text", filename="note.txt")
    assert exc.value.status_code == 503
    assert exc.value.detail == "documents.errors.malware_scan_unavailable"
