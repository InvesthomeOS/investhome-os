"""Sprint 5 — AI Index foundation tests (mocked Drive / extractors)."""

from __future__ import annotations

import io
import json
from collections.abc import Iterator
from datetime import UTC, datetime
from uuid import uuid4

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter
from sqlalchemy import select
from sqlalchemy.orm import Session

from investhome_api.db.session import get_db
from investhome_api.main import app
from investhome_api.config.settings import get_settings
from investhome_api.models.ai_index import AiDocument, AiDocumentStatus, AiDocumentType
from investhome_api.models.creative_studio_media import (
    CreativeStudioMediaAsset,
    MediaAssetSourceType,
    MediaAssetSyncStatus,
)
from investhome_api.models.project import Project, ProjectStatus, ProjectType
from investhome_api.models.project_drive import ProjectDriveMapping
from investhome_api.services.ai_index.extraction import extract_bytes
from investhome_api.services.ai_index.keywords import extract_keywords
from investhome_api.services.ai_index.normalizer import normalize_text
from investhome_api.services.ai_index.pipeline import (
    index_asset,
    index_special_drive_file,
    mark_inactive_for_asset,
)
from investhome_api.services.ai_index.summarizer import build_summary
from investhome_api.services.google_drive.errors import GoogleDriveApiError
from investhome_api.services.google_drive.provider import DriveFileMeta
from investhome_api.services.google_drive.scanner import DriveAssetScanner, upsert_project_drive_mapping

ROOT_ID = "root-folder"
PROJECT_FOLDER = "project-folder-ai"


def _meta(
    file_id: str,
    name: str,
    *,
    mime: str = "text/plain",
    parent: str = PROJECT_FOLDER,
    size: int = 100,
    modified: datetime | None = None,
    md5: str | None = "chk-1",
    is_folder: bool = False,
) -> DriveFileMeta:
    return DriveFileMeta(
        id=file_id,
        name=name,
        mime_type="application/vnd.google-apps.folder" if is_folder else mime,
        size=None if is_folder else size,
        created_at=datetime(2024, 1, 1, tzinfo=UTC),
        modified_at=modified or datetime(2024, 6, 1, tzinfo=UTC),
        md5_checksum=None if is_folder else md5,
        parent_ids=(parent,),
        web_view_link=f"https://drive.example/{file_id}",
        thumbnail_link=None,
    )


class FakeDriveProvider:
    def __init__(
        self,
        *,
        root_id: str = ROOT_ID,
        children: dict[str, list[DriveFileMeta]] | None = None,
        files: dict[str, DriveFileMeta] | None = None,
        downloads: dict[str, bytes] | None = None,
    ) -> None:
        self._root_id = root_id
        self.children = children or {}
        self.files = files or {}
        self.downloads = downloads or {}
        for items in self.children.values():
            for item in items:
                self.files.setdefault(item.id, item)
        self.files.setdefault(root_id, _meta(root_id, "Root", parent="", is_folder=True))

    @property
    def root_folder_id(self) -> str:
        return self._root_id

    def get_file(self, file_id: str) -> DriveFileMeta:
        if file_id not in self.files:
            raise GoogleDriveApiError(f"not found: {file_id}", code="drive_not_found")
        return self.files[file_id]

    def list_children(self, folder_id: str, *, page_size: int = 100) -> Iterator[DriveFileMeta]:
        yield from self.children.get(folder_id, [])

    def is_descendant_of_root(self, folder_id: str) -> bool:
        if folder_id == self._root_id:
            return True
        current = folder_id
        seen: set[str] = set()
        while current and current not in seen:
            seen.add(current)
            if current == self._root_id:
                return True
            meta = self.files.get(current)
            if meta is None or not meta.parent_ids:
                return False
            current = meta.parent_ids[0]
        return False

    def download_bytes(self, file_id: str, *, max_bytes: int | None = None) -> bytes:
        if file_id not in self.downloads:
            raise GoogleDriveApiError("download missing", code="drive_not_found")
        data = self.downloads[file_id]
        if max_bytes is not None and len(data) > max_bytes:
            raise GoogleDriveApiError("too large", code="drive_download_too_large")
        return data

    def get_start_page_token(self) -> str:
        return "token-1"

    def list_changes(self, page_token: str, *, page_size: int = 100):
        raise NotImplementedError


def _db() -> Session:
    return next(app.dependency_overrides[get_db]())


def _create_project(db: Session) -> Project:
    project = Project(
        id=uuid4(),
        project_code=f"PRJ-AI-{uuid4().hex[:8]}",
        project_name="AI Index Project",
        project_type=ProjectType.RESIDENTIAL,
        project_status=ProjectStatus.CONSTRUCTION,
    )
    db.add(project)
    db.flush()
    return project


def _docx_bytes(text: str) -> bytes:
    doc = DocxDocument()
    doc.add_heading("Project Brief", level=1)
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_bytes(text: str) -> bytes:
    # Minimal PDF with a text content stream (pypdf can re-read embedded text)
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject, NumberObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)
    # Simple text stream — extract_text may be empty on blank; add content
    stream = DecodedStreamObject()
    # Use a basic PDF text drawing command
    content = f"BT /F1 12 Tf 50 150 Td ({text}) Tj ET".encode("latin-1", errors="replace")
    stream.set_data(content)
    page[NameObject("/Contents")] = stream
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _standard_tree(downloads: dict[str, bytes]) -> FakeDriveProvider:
    info = _meta("folder-info", "00_PROJECT_INFO", is_folder=True, parent=PROJECT_FOLDER, md5=None)
    docs = _meta("folder-docs", "09_DOCUMENTS", is_folder=True, parent=PROJECT_FOLDER, md5=None)
    catalog = _meta("folder-catalog", "07_CATALOG", is_folder=True, parent=PROJECT_FOLDER, md5=None)

    readme = _meta("file-readme", "README.md", mime="text/markdown", parent=PROJECT_FOLDER, md5="rmd1")
    meta_json = _meta(
        "file-meta", "metadata.json", mime="application/json", parent="folder-info", md5="meta1"
    )
    notes = _meta("file-txt", "notes.txt", mime="text/plain", parent="folder-info", md5="txt1")
    brief = _meta(
        "file-docx",
        "brief.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        parent="folder-info",
        md5="docx1",
    )
    flyer = _meta("file-pdf", "flyer.pdf", mime="application/pdf", parent="folder-catalog", md5="pdf1")
    legal = _meta("file-legal", "contract.pdf", mime="application/pdf", parent="folder-docs", md5="leg1")
    image = _meta("file-img", "hero.png", mime="image/png", parent=PROJECT_FOLDER, md5="img1")

    children = {
        ROOT_ID: [_meta(PROJECT_FOLDER, "Project A", is_folder=True, parent=ROOT_ID, md5=None)],
        PROJECT_FOLDER: [info, docs, catalog, readme, image],
        "folder-info": [meta_json, notes, brief],
        "folder-docs": [legal],
        "folder-catalog": [flyer],
    }
    files = {m.id: m for items in children.values() for m in items}
    files[PROJECT_FOLDER] = _meta(PROJECT_FOLDER, "Project A", is_folder=True, parent=ROOT_ID, md5=None)
    return FakeDriveProvider(children=children, files=files, downloads=downloads)


@pytest.fixture
def db_session(client) -> Session:
    return _db()


def test_normalize_and_summary_keywords() -> None:
    raw = "Hello   world.\r\n\r\n# Title One\n\nThis is a substantial first sentence about housing. Another sentence follows here."
    text = normalize_text(raw)
    assert "\r" not in text
    assert "  " not in text
    summary = build_summary(text)
    assert "Title One" in summary
    keywords = extract_keywords(
        "housing housing housing apartment apartment development development development market"
    )
    assert keywords[0] == "housing"
    assert "development" in keywords


def test_extract_json_and_txt_and_docx() -> None:
    js = extract_bytes(json.dumps({"name": "Tower", "builders": ["Acme"]}).encode(), "metadata.json")
    assert js.document_type == AiDocumentType.METADATA.value
    assert "Tower" in js.text
    assert js.structured and js.structured["name"] == "Tower"

    md = extract_bytes(b"# Hello\n\nProject context here.", "README.md")
    assert md.document_type == AiDocumentType.README.value

    txt = extract_bytes(b"plain notes", "notes.txt")
    assert txt.document_type == AiDocumentType.TXT.value

    docx = extract_bytes(_docx_bytes("Luxury residences near the park."), "brief.docx")
    assert docx.document_type == AiDocumentType.DOCX.value
    assert "Luxury" in docx.text

    skipped = extract_bytes(b"\x89PNG", "hero.png")
    assert skipped.skipped is True


def test_readme_and_metadata_indexed(db_session: Session) -> None:
    db = db_session
    project = _create_project(db)
    downloads = {
        "file-readme": b"# Seaside Residences\n\nFull project context for agents.",
        "file-meta": json.dumps(
            {"name": "Seaside", "builders": ["BuildCo"], "city": "Izmir"}
        ).encode(),
        "file-txt": b"Construction notes: phase 1 foundation complete.",
        "file-docx": _docx_bytes("Marketing brief for Seaside residences near the marina."),
        "file-pdf": b"%PDF-1.4 fake",  # may extract empty — still indexes
        "file-legal": b"%PDF-1.4 legal",
        "file-img": b"\x89PNG",
    }
    provider = _standard_tree(downloads)
    upsert_project_drive_mapping(
        db, project_id=project.id, drive_folder_id=PROJECT_FOLDER, provider=provider
    )
    scanner = DriveAssetScanner(db, provider, project_id=project.id)
    summary = scanner.sync(dry_run=False)
    db.commit()

    assert summary.created >= 1

    readme_doc = db.scalar(
        select(AiDocument).where(
            AiDocument.project_id == project.id,
            AiDocument.drive_file_id == "file-readme",
        )
    )
    assert readme_doc is not None
    assert readme_doc.document_type == AiDocumentType.README.value
    assert readme_doc.index_status == AiDocumentStatus.READY.value
    assert "Seaside Residences" in (readme_doc.extracted_text or "")
    assert readme_doc.summary
    assert readme_doc.keywords
    assert readme_doc.asset_id is None  # README is not a Media Library asset

    meta_doc = db.scalar(
        select(AiDocument).where(
            AiDocument.project_id == project.id,
            AiDocument.drive_file_id == "file-meta",
        )
    )
    assert meta_doc is not None
    assert meta_doc.document_type == AiDocumentType.METADATA.value
    assert meta_doc.builders == ["BuildCo"]
    assert "Izmir" in (meta_doc.extracted_text or "")


def test_txt_docx_json_pdf_and_legal_meta(db_session: Session) -> None:
    db = db_session
    project = _create_project(db)
    downloads = {
        "file-readme": b"# R\n",
        "file-meta": b'{"index": true}',
        "file-txt": b"Location notes about the waterfront district amenities.",
        "file-docx": _docx_bytes("Detailed catalog notes for unit types and finishes."),
        "file-pdf": _pdf_bytes("Catalog flyer text layer content for buyers."),
        "file-legal": b"%PDF-1.4 contract body should not be fully indexed as legal",
        "file-img": b"img",
    }
    provider = _standard_tree(downloads)
    upsert_project_drive_mapping(
        db, project_id=project.id, drive_folder_id=PROJECT_FOLDER, provider=provider
    )
    DriveAssetScanner(db, provider, project_id=project.id).sync(dry_run=False)
    db.commit()

    txt_asset = db.scalar(
        select(CreativeStudioMediaAsset).where(CreativeStudioMediaAsset.external_file_id == "file-txt")
    )
    assert txt_asset is not None
    txt_doc = db.scalar(select(AiDocument).where(AiDocument.asset_id == txt_asset.id))
    assert txt_doc is not None
    assert txt_doc.document_type == AiDocumentType.TXT.value
    assert "waterfront" in (txt_doc.extracted_text or "").lower()

    docx_asset = db.scalar(
        select(CreativeStudioMediaAsset).where(CreativeStudioMediaAsset.external_file_id == "file-docx")
    )
    docx_doc = db.scalar(select(AiDocument).where(AiDocument.asset_id == docx_asset.id))
    assert docx_doc is not None
    assert docx_doc.document_type == AiDocumentType.DOCX.value

    legal_asset = db.scalar(
        select(CreativeStudioMediaAsset).where(CreativeStudioMediaAsset.external_file_id == "file-legal")
    )
    legal_doc = db.scalar(select(AiDocument).where(AiDocument.asset_id == legal_asset.id))
    assert legal_doc is not None
    assert legal_doc.document_type == AiDocumentType.LEGAL_META.value
    assert legal_doc.metadata_json and legal_doc.metadata_json.get("legal_meta_only") is True
    # Should not contain raw PDF binary payload
    assert "%PDF" not in (legal_doc.extracted_text or "")


def test_checksum_skips_unchanged(db_session: Session) -> None:
    db = db_session
    project = _create_project(db)
    provider = FakeDriveProvider(
        children={
            ROOT_ID: [_meta(PROJECT_FOLDER, "P", is_folder=True, parent=ROOT_ID, md5=None)],
            PROJECT_FOLDER: [
                _meta("folder-info", "00_PROJECT_INFO", is_folder=True, parent=PROJECT_FOLDER, md5=None),
                _meta("file-readme", "README.md", mime="text/markdown", md5="same"),
            ],
            "folder-info": [
                _meta("file-txt", "notes.txt", parent="folder-info", md5="same-txt"),
            ],
        },
        downloads={
            "file-readme": b"# Same\n\nUnchanged project context paragraph one.",
            "file-txt": b"Unchanged notes about the residential tower amenities.",
        },
    )
    # Fix files map parents
    provider.files["folder-info"] = _meta(
        "folder-info", "00_PROJECT_INFO", is_folder=True, parent=PROJECT_FOLDER, md5=None
    )
    provider.files[PROJECT_FOLDER] = _meta(PROJECT_FOLDER, "P", is_folder=True, parent=ROOT_ID, md5=None)

    upsert_project_drive_mapping(
        db, project_id=project.id, drive_folder_id=PROJECT_FOLDER, provider=provider
    )
    DriveAssetScanner(db, provider, project_id=project.id).sync(dry_run=False)
    db.flush()

    doc = db.scalar(
        select(AiDocument).where(
            AiDocument.project_id == project.id, AiDocument.drive_file_id == "file-readme"
        )
    )
    assert doc is not None
    version_1 = doc.version
    indexed_at = doc.last_indexed_at

    # Second sync — unchanged checksum should skip rewrite
    DriveAssetScanner(db, provider, project_id=project.id).sync(dry_run=False)
    db.flush()
    db.refresh(doc)
    assert doc.version == version_1
    assert doc.checksum == "same"

    # Force reindex still works
    index_special_drive_file(
        db,
        project_id=project.id,
        drive_file_id="file-readme",
        filename="README.md",
        category=None,
        provider=provider,
        checksum="same",
        force=True,
    )
    db.flush()
    db.refresh(doc)
    assert doc.last_indexed_at is not None


def test_deleted_drive_file_marks_inactive(db_session: Session) -> None:
    db = db_session
    project = _create_project(db)
    provider = FakeDriveProvider(
        children={
            ROOT_ID: [_meta(PROJECT_FOLDER, "P", is_folder=True, parent=ROOT_ID, md5=None)],
            PROJECT_FOLDER: [
                _meta("folder-info", "00_PROJECT_INFO", is_folder=True, parent=PROJECT_FOLDER, md5=None),
            ],
            "folder-info": [
                _meta("file-txt", "notes.txt", parent="folder-info", md5="t1"),
            ],
        },
        downloads={"file-txt": b"Notes about construction schedule and deliveries."},
    )
    provider.files["folder-info"] = _meta(
        "folder-info", "00_PROJECT_INFO", is_folder=True, parent=PROJECT_FOLDER, md5=None
    )
    provider.files[PROJECT_FOLDER] = _meta(PROJECT_FOLDER, "P", is_folder=True, parent=ROOT_ID, md5=None)

    upsert_project_drive_mapping(
        db, project_id=project.id, drive_folder_id=PROJECT_FOLDER, provider=provider
    )
    DriveAssetScanner(db, provider, project_id=project.id).sync(dry_run=False)
    db.flush()

    asset = db.scalar(
        select(CreativeStudioMediaAsset).where(CreativeStudioMediaAsset.external_file_id == "file-txt")
    )
    doc = db.scalar(select(AiDocument).where(AiDocument.asset_id == asset.id))
    assert doc is not None
    assert doc.is_active is True

    # File removed from Drive tree
    provider.children["folder-info"] = []
    DriveAssetScanner(db, provider, project_id=project.id).sync(dry_run=False)
    db.flush()
    db.refresh(asset)
    db.refresh(doc)
    assert asset.sync_status == MediaAssetSyncStatus.MISSING.value
    assert doc.is_active is False
    assert doc.index_status == AiDocumentStatus.INACTIVE.value


def test_api_list_and_asset_ai_document(client, db_session: Session) -> None:
    db = db_session
    project = _create_project(db)
    asset = CreativeStudioMediaAsset(
        id=uuid4(),
        filename="notes.txt",
        content_type="text/plain",
        file_size=20,
        storage_provider="google_drive",
        storage_key="gdrive:x",
        linked_project_id=project.id,
        source_type=MediaAssetSourceType.GOOGLE_DRIVE.value,
        external_file_id="x",
        external_checksum="c1",
        sync_status=MediaAssetSyncStatus.ACTIVE.value,
        folder_category="00_PROJECT_INFO",
    )
    db.add(asset)
    db.flush()

    class _Prov(FakeDriveProvider):
        pass

    provider = _Prov(downloads={"x": b"API notes about the residential amenities nearby."})
    index_asset(db, asset.id, provider=provider)
    db.commit()

    listed = client.get(f"/projects/{project.id}/ai-index")
    assert listed.status_code == 200
    body = listed.json()
    assert body["total"] >= 1
    assert body["items"][0]["summary"]

    one = client.get(f"/creative-studio/media/assets/{asset.id}/ai-document")
    assert one.status_code == 200
    assert one.json()["asset_id"] == str(asset.id)
    assert one.json()["document_type"] == "txt"


def test_enqueue_asset_ai_index_waits_for_commit(db_session: Session, monkeypatch: pytest.MonkeyPatch) -> None:
    """Async enqueue must not fire until the sync transaction commits (Drive race fix)."""
    from investhome_api.services.ai_index import queue as ai_queue

    monkeypatch.setenv("AI_INDEX_PROCESSING_SYNC", "false")
    get_settings.cache_clear()

    dispatched: list[tuple[object, bool]] = []

    def _capture(asset_id, *, force: bool = False) -> None:
        dispatched.append((asset_id, force))

    monkeypatch.setattr(ai_queue, "_dispatch_asset_enqueue", _capture)

    asset_id = uuid4()
    ai_queue.enqueue_asset_ai_index(asset_id, force=True, db=db_session)
    assert dispatched == []

    db_session.commit()
    assert dispatched == [(asset_id, True)]

    # Restore sync mode for other autouse expectations within the same process
    monkeypatch.setenv("AI_INDEX_PROCESSING_SYNC", "true")
    get_settings.cache_clear()


def test_enqueue_asset_ai_index_skips_on_rollback(db_session: Session, monkeypatch: pytest.MonkeyPatch) -> None:
    from investhome_api.services.ai_index import queue as ai_queue

    monkeypatch.setenv("AI_INDEX_PROCESSING_SYNC", "false")
    get_settings.cache_clear()

    dispatched: list[object] = []
    monkeypatch.setattr(
        ai_queue, "_dispatch_asset_enqueue", lambda asset_id, *, force=False: dispatched.append(asset_id)
    )

    asset_id = uuid4()
    ai_queue.enqueue_asset_ai_index(asset_id, db=db_session)
    db_session.rollback()
    assert dispatched == []

    monkeypatch.setenv("AI_INDEX_PROCESSING_SYNC", "true")
    get_settings.cache_clear()
