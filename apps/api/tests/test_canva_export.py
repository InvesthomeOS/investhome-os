"""Canva Connect export — PNG/asset upload, create design, token refresh. No secrets in responses."""

from __future__ import annotations

import base64
import io
import json
from typing import Any
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from sqlalchemy.orm import Session

from investhome_api.config.settings import get_settings
from investhome_api.services import canva_layered_pptx as layered
from investhome_api.services import canva_oauth_service as canva
from investhome_api.services.storage.factory import get_storage_provider

EDIT_URL = "https://www.canva.com/api/design/test-edit-token/edit"
DOWNLOAD_URL = "https://export-download.canva.com/preview.png"


class _FakeResponse:
    def __init__(
        self,
        status_code: int,
        payload: dict[str, Any] | None = None,
        content: bytes | None = None,
    ):
        self.status_code = status_code
        self._payload = payload or {}
        self.content = content if content is not None else b""

    def json(self) -> dict[str, Any]:
        return self._payload


class _FakeCanvaHttp:
    def __init__(self) -> None:
        self.calls: list[tuple[str, str]] = []
        self.refresh_count = 0
        self.upload_count = 0
        self.poll_count = 0
        self.design_count = 0
        self.import_count = 0
        self.import_poll_count = 0
        self.export_count = 0
        self.export_poll_count = 0
        self.download_count = 0
        self.last_import_body: bytes | None = None
        self.last_export_body: dict[str, Any] | None = None
        self.seen_authorization: list[str] = []
        self.poll_in_progress_once = False
        self.export_poll_in_progress_once = False
        self.fail_first_upload_auth = False
        self.fail_import = False
        self.fail_export = False
        self.preview_png = _png_bytes(32, 32)
        self._upload_auth_failures = 0

    def __enter__(self) -> _FakeCanvaHttp:
        return self

    def __exit__(self, *args: object) -> None:
        return None

    def _record(self, method: str, url: str, headers: dict | None) -> None:
        self.calls.append((method, url))
        auth = (headers or {}).get("Authorization")
        if auth:
            self.seen_authorization.append(str(auth))

    def _export_job(self, *, success: bool) -> dict[str, Any]:
        job: dict[str, Any] = {"id": "job-export-1", "status": "success" if success else "in_progress"}
        if success:
            job["urls"] = [DOWNLOAD_URL]
        return job

    def post(self, url: str, headers: dict | None = None, data: Any = None, content: Any = None, json: Any = None):
        self._record("POST", url, headers)
        if url == canva.CANVA_TOKEN_URL:
            form = data or {}
            assert form.get("grant_type") == "refresh_token"
            assert form.get("refresh_token")
            self.refresh_count += 1
            return _FakeResponse(
                200,
                {
                    "access_token": "access-refreshed",
                    "refresh_token": "refresh-rotated",
                    "token_type": "Bearer",
                    "expires_in": 3600,
                },
            )
        if url.endswith("/asset-uploads"):
            if self.fail_first_upload_auth and self._upload_auth_failures == 0:
                self._upload_auth_failures += 1
                return _FakeResponse(401, {"code": "unauthorized"})
            self.upload_count += 1
            status = "in_progress" if self.poll_in_progress_once else "success"
            job: dict[str, Any] = {"id": "job-upload-1", "status": status}
            if status == "success":
                job["asset"] = {"id": "MsdCanvaAsset", "type": "image", "name": "Social post"}
            return _FakeResponse(200, {"job": job})
        if url.rstrip("/").endswith("/imports"):
            self.import_count += 1
            self.last_import_body = content if isinstance(content, (bytes, bytearray)) else None
            if self.fail_import:
                return _FakeResponse(400, {"code": "invalid_file"})
            return _FakeResponse(
                200,
                {
                    "job": {
                        "id": "job-import-1",
                        "status": "success",
                        "result": {
                            "designs": [
                                {
                                    "id": "DAGeditable1",
                                    "urls": {
                                        "edit_url": EDIT_URL,
                                        "view_url": "https://www.canva.com/api/design/test-edit-token/view",
                                    },
                                }
                            ]
                        },
                    }
                },
            )
        if url.rstrip("/").endswith("/exports"):
            self.export_count += 1
            self.last_export_body = json if isinstance(json, dict) else None
            if self.fail_export:
                return _FakeResponse(403, {"code": "permission_denied"})
            assert json and json.get("design_id")
            assert json.get("format", {}).get("type") == "png"
            success = not self.export_poll_in_progress_once
            return _FakeResponse(200, {"job": self._export_job(success=success)})
        if url.endswith("/designs"):
            self.design_count += 1
            assert json and json.get("asset_id") == "MsdCanvaAsset"
            return _FakeResponse(
                200,
                {
                    "design": {
                        "id": "DAFdesign1",
                        "title": json.get("title"),
                        "urls": {
                            "edit_url": EDIT_URL,
                            "view_url": "https://www.canva.com/api/design/test-edit-token/view",
                        },
                    }
                },
            )
        return _FakeResponse(404, {"code": "not_found"})

    def get(self, url: str, headers: dict | None = None, follow_redirects: bool = False, **_kwargs: Any):
        self._record("GET", url, headers)
        _ = follow_redirects
        if url.startswith(DOWNLOAD_URL) or "export-download.canva.com" in url:
            self.download_count += 1
            return _FakeResponse(200, content=self.preview_png)
        if "/exports/" in url:
            self.export_poll_count += 1
            return _FakeResponse(200, {"job": self._export_job(success=True)})
        if "/imports/" in url:
            self.import_poll_count += 1
            return _FakeResponse(
                200,
                {
                    "job": {
                        "id": "job-import-1",
                        "status": "success",
                        "result": {
                            "designs": [
                                {
                                    "id": "DAGeditable1",
                                    "urls": {
                                        "edit_url": EDIT_URL,
                                        "view_url": "https://www.canva.com/api/design/test-edit-token/view",
                                    },
                                }
                            ]
                        },
                    }
                },
            )
        self.poll_count += 1
        return _FakeResponse(
            200,
            {
                "job": {
                    "id": "job-upload-1",
                    "status": "success",
                    "asset": {"id": "MsdCanvaAsset", "type": "image", "name": "Social post"},
                }
            },
        )


@pytest.fixture(autouse=True)
def _canva_env(monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setenv("CANVA_CLIENT_ID", "test-canva-client-id")
    monkeypatch.setenv("CANVA_CLIENT_SECRET", "test-canva-client-secret")
    get_settings.cache_clear()
    yield
    get_settings.cache_clear()


def _png_bytes(width: int = 64, height: int = 64) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), (12, 80, 160)).save(buf, format="PNG")
    return buf.getvalue()


def _store_tokens(db: Session, *, access: str = "access-live", refresh: str = "refresh-live", expires_in: int = 3600) -> None:
    canva.store_tokens(
        db,
        {
            "access_token": access,
            "refresh_token": refresh,
            "token_type": "Bearer",
            "expires_in": expires_in,
            "scope": " ".join(canva.CANVA_SCOPES),
        },
    )
    db.commit()


def _secret_leak(payload: str) -> None:
    lowered = payload.lower()
    assert "access-live" not in lowered
    assert "refresh-live" not in lowered
    assert "access-refreshed" not in lowered
    assert "refresh-rotated" not in lowered
    assert "test-canva-client-secret" not in lowered
    assert "bearer access" not in lowered


def test_export_requires_canva_connection(client: TestClient) -> None:
    png = _png_bytes()
    response = client.post(
        "/platform/integrations/canva/export",
        files={"file": ("post.png", io.BytesIO(png), "image/png")},
        data={"title": "Temple post"},
    )
    assert response.status_code == 409
    assert response.json()["detail"] == "canva_not_connected"
    _secret_leak(response.text)


def test_export_png_uploads_and_returns_edit_url(client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch) -> None:
    fake = _FakeCanvaHttp()
    monkeypatch.setattr(canva, "http_client", lambda: fake)
    monkeypatch.setattr("investhome_api.services.canva_export_service._sleep", lambda _s: None)
    _store_tokens(db)

    png = _png_bytes(1080, 1080)
    response = client.post(
        "/platform/integrations/canva/export",
        files={"file": ("temple.png", io.BytesIO(png), "image/png")},
        data={"title": "Temple Instagram", "width": "1080", "height": "1080"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["edit_url"] == EDIT_URL
    assert body["design_id"] == "DAFdesign1"
    decoded = base64.b64decode(body["preview_png_base64"])
    assert decoded[:8] == b"\x89PNG\r\n\x1a\n"
    assert fake.upload_count == 1
    assert fake.design_count == 1
    assert fake.export_count == 1
    assert fake.download_count == 1
    assert fake.refresh_count == 0
    _secret_leak(response.text)
    assert "access_token" not in response.text.lower()
    assert "refresh_token" not in response.text.lower()


def test_export_polls_in_progress_upload_job(client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch) -> None:
    fake = _FakeCanvaHttp()
    fake.poll_in_progress_once = True
    monkeypatch.setattr(canva, "http_client", lambda: fake)
    monkeypatch.setattr("investhome_api.services.canva_export_service._sleep", lambda _s: None)
    _store_tokens(db)

    response = client.post(
        "/platform/integrations/canva/export",
        files={"file": ("post.png", io.BytesIO(_png_bytes()), "image/png")},
        data={"title": "Poll me"},
    )
    assert response.status_code == 200, response.text
    assert fake.poll_count >= 1
    assert response.json()["edit_url"] == EDIT_URL


def test_export_refreshes_expired_access_token(client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch) -> None:
    fake = _FakeCanvaHttp()
    monkeypatch.setattr(canva, "http_client", lambda: fake)
    monkeypatch.setattr("investhome_api.services.canva_export_service._sleep", lambda _s: None)
    _store_tokens(db, access="access-expired", refresh="refresh-live", expires_in=-30)

    response = client.post(
        "/platform/integrations/canva/export",
        files={"file": ("post.png", io.BytesIO(_png_bytes()), "image/png")},
        data={"title": "Refresh path"},
    )
    assert response.status_code == 200, response.text
    assert fake.refresh_count == 1
    assert any(auth.endswith("access-refreshed") for auth in fake.seen_authorization)
    stored = canva.get_stored_tokens(db)
    assert stored is not None
    assert stored["access_token"] == "access-refreshed"
    assert stored["refresh_token"] == "refresh-rotated"
    _secret_leak(response.text)


def test_export_retries_upload_after_401(client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch) -> None:
    fake = _FakeCanvaHttp()
    fake.fail_first_upload_auth = True
    monkeypatch.setattr(canva, "http_client", lambda: fake)
    monkeypatch.setattr("investhome_api.services.canva_export_service._sleep", lambda _s: None)
    _store_tokens(db, access="access-stale", refresh="refresh-live", expires_in=3600)

    response = client.post(
        "/platform/integrations/canva/export",
        files={"file": ("post.png", io.BytesIO(_png_bytes()), "image/png")},
        data={"title": "401 retry"},
    )
    assert response.status_code == 200, response.text
    assert fake.refresh_count == 1
    assert fake.upload_count == 1
    _secret_leak(response.text)


def test_export_rejects_non_png(client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch) -> None:
    fake = _FakeCanvaHttp()
    monkeypatch.setattr(canva, "http_client", lambda: fake)
    _store_tokens(db)
    response = client.post(
        "/platform/integrations/canva/export",
        files={"file": ("note.txt", io.BytesIO(b"not-an-image"), "text/plain")},
        data={"title": "bad"},
    )
    assert response.status_code == 400
    assert response.json()["detail"] == "invalid_png"
    assert fake.upload_count == 0


def test_export_media_asset_enforces_project_isolation(
    client: TestClient,
    db: Session,
    tmp_path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from investhome_api.models.project import Project, ProjectStatus, ProjectType

    monkeypatch.setenv("DOCUMENT_STORAGE_ROOT", str(tmp_path / "media-storage"))
    get_settings.cache_clear()
    get_storage_provider.cache_clear()

    fake = _FakeCanvaHttp()
    monkeypatch.setattr(canva, "http_client", lambda: fake)
    monkeypatch.setattr("investhome_api.services.canva_export_service._sleep", lambda _s: None)
    _store_tokens(db)

    project_a = Project(
        id=uuid4(),
        project_code=f"PRJ-A-{uuid4().hex[:6]}",
        project_name="Temple",
        project_type=ProjectType.RESIDENTIAL,
        project_status=ProjectStatus.CONSTRUCTION,
    )
    project_b = Project(
        id=uuid4(),
        project_code=f"PRJ-B-{uuid4().hex[:6]}",
        project_name="Other",
        project_type=ProjectType.RESIDENTIAL,
        project_status=ProjectStatus.CONSTRUCTION,
    )
    db.add(project_a)
    db.add(project_b)
    db.commit()

    png = _png_bytes()
    upload = client.post(
        "/creative-studio/media/upload",
        files={"file": ("hero.png", io.BytesIO(png), "image/png")},
        data={"linked_project_id": str(project_a.id), "tags": "hero"},
    )
    assert upload.status_code == 201, upload.text
    asset_id = upload.json()["id"]

    denied = client.post(
        "/platform/integrations/canva/export",
        data={
            "media_asset_id": asset_id,
            "linked_project_id": str(project_b.id),
            "title": "cross project",
        },
    )
    assert denied.status_code == 404
    assert fake.upload_count == 0
    _secret_leak(denied.text)

    allowed = client.post(
        "/platform/integrations/canva/export",
        data={
            "media_asset_id": asset_id,
            "linked_project_id": str(project_a.id),
            "title": "same project",
        },
    )
    assert allowed.status_code == 200, allowed.text
    assert allowed.json()["edit_url"] == EDIT_URL
    assert fake.upload_count == 1


def test_export_media_asset_requires_project_id(client: TestClient, db: Session) -> None:
    _store_tokens(db)
    response = client.post(
        "/platform/integrations/canva/export",
        data={"media_asset_id": str(uuid4()), "title": "no project"},
    )
    assert response.status_code == 400
    assert response.json()["detail"] == "linked_project_id_required"


def test_get_valid_access_token_refreshes(db: Session, monkeypatch: pytest.MonkeyPatch) -> None:
    fake = _FakeCanvaHttp()
    monkeypatch.setattr(canva, "http_client", lambda: fake)
    _store_tokens(db, access="old", refresh="refresh-live", expires_in=-5)
    token = canva.get_valid_access_token(db)
    assert token == "access-refreshed"
    assert fake.refresh_count == 1
    stored = canva.get_stored_tokens(db)
    assert stored is not None
    assert stored["refresh_token"] == "refresh-rotated"


def _layered_payload() -> str:
    return json.dumps(
        {
            "width": 1080,
            "height": 1080,
            "brand_logo": True,
            "cover_asset_key": "cover.png",
            "elements": [
                {
                    "id": "h1",
                    "type": "TEXT",
                    "role": "headline",
                    "content": "Temple headline",
                    "x": 80,
                    "y": 200,
                    "width": 920,
                    "height": 80,
                    "zIndex": 2,
                    "fontSize": 48,
                    "fontWeight": "bold",
                    "align": "center",
                    "color": "#ffffff",
                },
                {
                    "id": "b1",
                    "type": "TEXT",
                    "role": "body",
                    "content": "From 4.2M AED",
                    "x": 80,
                    "y": 300,
                    "width": 920,
                    "height": 50,
                    "zIndex": 3,
                    "fontSize": 24,
                    "fontWeight": "normal",
                    "align": "center",
                    "color": "#ffffff",
                },
                {
                    "id": "c1",
                    "type": "BUTTON",
                    "label": "Book a tour",
                    "x": 340,
                    "y": 900,
                    "width": 400,
                    "height": 48,
                    "zIndex": 4,
                    "backgroundColor": "#ffffff",
                    "textColor": "#111827",
                },
                {
                    "id": "m1",
                    "type": "METRIC_GROUP",
                    "x": 80,
                    "y": 700,
                    "width": 920,
                    "height": 80,
                    "zIndex": 5,
                    "color": "#ffffff",
                    "layout": "horizontal",
                    "metrics": [{"display_value": "12%", "label": "Target yield"}],
                },
                {
                    "id": "logo1",
                    "type": "IMAGE",
                    "role": "logo",
                    "asset_key": "logo.png",
                    "x": 40,
                    "y": 40,
                    "width": 120,
                    "height": 48,
                    "zIndex": 6,
                },
            ],
        }
    )


def test_pptx_layers_are_separate_text_and_pictures() -> None:
    png = _png_bytes(64, 64)
    parsed = layered.parse_layers_json(_layered_payload())
    assert parsed is not None
    pptx = layered.build_pptx_from_layers(parsed, {"cover.png": png, "logo.png": png})
    assert pptx is not None
    assert pptx[:2] == b"PK"
    prs = Presentation(io.BytesIO(pptx))
    slide = prs.slides[0]
    texts = [
        shape.text_frame.text
        for shape in slide.shapes
        if shape.has_text_frame and shape.text_frame.text.strip()
    ]
    pictures = [shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE]
    joined = "\n".join(texts)
    assert "Temple headline" in joined
    assert "From 4.2M AED" in joined
    assert "Book a tour" in joined
    assert "12%" in joined
    assert "Target yield" in joined
    assert any(text.strip() == "IH" for text in texts)
    assert len(pictures) >= 2


def test_export_layers_uses_design_import(
    client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake = _FakeCanvaHttp()
    monkeypatch.setattr(canva, "http_client", lambda: fake)
    monkeypatch.setattr("investhome_api.services.canva_export_service._sleep", lambda _s: None)
    _store_tokens(db)

    png = _png_bytes(1080, 1080)
    response = client.post(
        "/platform/integrations/canva/export",
        files=[
            ("file", ("temple.png", io.BytesIO(png), "image/png")),
            ("layer_images", ("cover.png", io.BytesIO(png), "image/png")),
            ("layer_images", ("logo.png", io.BytesIO(_png_bytes(32, 32)), "image/png")),
        ],
        data={
            "title": "Temple Instagram",
            "width": "1080",
            "height": "1080",
            "layers": _layered_payload(),
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["edit_url"] == EDIT_URL
    assert body["design_id"] == "DAGeditable1"
    assert body["transfer_mode"] == "editable"
    decoded = base64.b64decode(body["preview_png_base64"])
    assert decoded[:8] == b"\x89PNG\r\n\x1a\n"
    assert fake.import_count == 1
    assert fake.export_count == 1
    assert fake.download_count == 1
    assert fake.upload_count == 0
    assert fake.design_count == 0
    assert fake.last_import_body is not None
    assert fake.last_import_body[:2] == b"PK"
    _secret_leak(response.text)


def test_export_layers_fall_back_to_png_when_import_fails(
    client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake = _FakeCanvaHttp()
    fake.fail_import = True
    monkeypatch.setattr(canva, "http_client", lambda: fake)
    monkeypatch.setattr("investhome_api.services.canva_export_service._sleep", lambda _s: None)
    _store_tokens(db)

    png = _png_bytes(64, 64)
    response = client.post(
        "/platform/integrations/canva/export",
        files=[
            ("file", ("post.png", io.BytesIO(png), "image/png")),
            ("layer_images", ("cover.png", io.BytesIO(png), "image/png")),
        ],
        data={"title": "Fallback", "layers": _layered_payload()},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["edit_url"] == EDIT_URL
    assert body["transfer_mode"] == "png"
    assert body["preview_png_base64"]
    assert fake.import_count == 1
    assert fake.upload_count == 1
    assert fake.design_count == 1
    assert fake.export_count == 1
    _secret_leak(response.text)


def test_export_returns_preview_png_after_polling_export_job(
    client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake = _FakeCanvaHttp()
    fake.export_poll_in_progress_once = True
    monkeypatch.setattr(canva, "http_client", lambda: fake)
    monkeypatch.setattr("investhome_api.services.canva_export_service._sleep", lambda _s: None)
    _store_tokens(db)

    response = client.post(
        "/platform/integrations/canva/export",
        files={"file": ("post.png", io.BytesIO(_png_bytes()), "image/png")},
        data={"title": "Poll export"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["edit_url"] == EDIT_URL
    decoded = base64.b64decode(body["preview_png_base64"])
    assert decoded[:8] == b"\x89PNG\r\n\x1a\n"
    assert fake.export_count == 1
    assert fake.export_poll_count >= 1
    assert fake.download_count == 1
    assert "access_token" not in response.text.lower()
    _secret_leak(response.text)


def test_export_keeps_edit_url_when_canva_png_render_fails(
    client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    fake = _FakeCanvaHttp()
    fake.fail_export = True
    monkeypatch.setattr(canva, "http_client", lambda: fake)
    monkeypatch.setattr("investhome_api.services.canva_export_service._sleep", lambda _s: None)
    _store_tokens(db)

    response = client.post(
        "/platform/integrations/canva/export",
        files={"file": ("post.png", io.BytesIO(_png_bytes()), "image/png")},
        data={"title": "Render fail"},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["edit_url"] == EDIT_URL
    assert body["design_id"] == "DAFdesign1"
    assert body.get("preview_png_base64") in (None, "")
    assert fake.export_count == 1
    assert fake.download_count == 0
    _secret_leak(response.text)
