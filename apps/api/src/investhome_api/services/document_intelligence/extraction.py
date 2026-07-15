"""Text extraction adapters for supported document formats."""

from __future__ import annotations

import csv
import io
import re
from typing import BinaryIO

import chardet
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from investhome_api.services.document_intelligence.types import (
    ExtractionResult,
    SourceReference,
    TextSegment,
)

MIN_PAGE_TEXT_CHARS = 40
TEXT_EXTRACTABLE = frozenset({"pdf", "docx", "xlsx", "pptx", "csv", "txt"})
IMAGE_EXTENSIONS = frozenset({"jpg", "jpeg", "png", "webp"})


def is_text_extractable(extension: str) -> bool:
    return extension.lower().lstrip(".") in TEXT_EXTRACTABLE


def is_image_extension(extension: str) -> bool:
    return extension.lower().lstrip(".") in IMAGE_EXTENSIONS


def is_processable(extension: str) -> bool:
    ext = extension.lower().lstrip(".")
    return ext in TEXT_EXTRACTABLE or ext in IMAGE_EXTENSIONS


def extract_text(content: bytes, extension: str) -> ExtractionResult:
    ext = extension.lower().lstrip(".")
    if ext == "pdf":
        return _extract_pdf(content)
    if ext == "docx":
        return _extract_docx(content)
    if ext == "xlsx":
        return _extract_xlsx(content)
    if ext == "pptx":
        return _extract_pptx(content)
    if ext == "csv":
        return _extract_csv(content)
    if ext == "txt":
        return _extract_txt(content)
    if ext in IMAGE_EXTENSIONS:
        return ExtractionResult(segments=[], method="image_pending_ocr")
    return ExtractionResult(segments=[], method="unsupported")


def _extract_pdf(content: bytes) -> ExtractionResult:
    reader = PdfReader(io.BytesIO(content))
    segments: list[TextSegment] = []
    for index, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        needs_ocr = len(re.sub(r"\s+", "", text)) < MIN_PAGE_TEXT_CHARS
        segments.append(
            TextSegment(
                text=text,
                source=SourceReference(page=index),
                needs_ocr=needs_ocr,
            )
        )
    result = ExtractionResult(
        segments=segments,
        method="pdf_embedded",
        page_count=len(reader.pages),
    )
    result.build_full_text()
    return result


def _extract_docx(content: bytes) -> ExtractionResult:
    doc = DocxDocument(io.BytesIO(content))
    segments: list[TextSegment] = []
    for index, para in enumerate(doc.paragraphs, start=1):
        text = para.text.strip()
        if text:
            style = para.style.name if para.style else "paragraph"
            segments.append(
                TextSegment(
                    text=text,
                    source=SourceReference(section=f"paragraph:{index} ({style})"),
                )
            )
    for t_index, table in enumerate(doc.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            segments.append(
                TextSegment(
                    text="\n".join(rows),
                    source=SourceReference(section=f"table:{t_index}"),
                )
            )
    result = ExtractionResult(segments=segments, method="docx")
    result.build_full_text()
    return result


def _extract_xlsx(content: bytes) -> ExtractionResult:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    segments: list[TextSegment] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows: list[str] = []
        for row in ws.iter_rows(values_only=True):
            values = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if values:
                rows.append(" | ".join(values))
        if rows:
            segments.append(
                TextSegment(
                    text="\n".join(rows[:500]),
                    source=SourceReference(sheet=sheet_name),
                )
            )
    wb.close()
    result = ExtractionResult(segments=segments, method="xlsx")
    result.build_full_text()
    return result


def _extract_pptx(content: bytes) -> ExtractionResult:
    prs = Presentation(io.BytesIO(content))
    segments: list[TextSegment] = []
    for index, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                texts.append(f"[notes] {notes}")
        if texts:
            segments.append(
                TextSegment(
                    text="\n".join(texts),
                    source=SourceReference(slide=index),
                )
            )
    result = ExtractionResult(segments=segments, method="pptx", page_count=len(prs.slides))
    result.build_full_text()
    return result


def _extract_csv(content: bytes) -> ExtractionResult:
    detected = chardet.detect(content[:10000])
    encoding = detected.get("encoding") or "utf-8"
    text = content.decode(encoding, errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [" | ".join(row) for row in reader if any(cell.strip() for cell in row)]
    limited = "\n".join(rows[:1000])
    segments = [TextSegment(text=limited, source=SourceReference(section="csv"))]
    result = ExtractionResult(segments=segments, method="csv")
    result.build_full_text()
    return result


def _extract_txt(content: bytes) -> ExtractionResult:
    detected = chardet.detect(content[:10000])
    encoding = detected.get("encoding") or "utf-8"
    text = content.decode(encoding, errors="replace")[:500_000]
    segments = [TextSegment(text=text, source=SourceReference(section="text"))]
    result = ExtractionResult(segments=segments, method="txt")
    result.full_text = text
    return result


def read_binary(stream: BinaryIO) -> bytes:
    return stream.read()
