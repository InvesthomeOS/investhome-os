"""Build a PPTX from SMB layers for Canva Design Import.

Canva Connect has no JSON "create elements" API. Official path with current
scopes (`design:content:write`) is POST /rest/v1/imports. PPTX is a documented
import type; Canva breaks text frames and pictures into separate editable
elements (see Design Import + Canva Help PDF/PPTX import).
"""

from __future__ import annotations

import contextlib
import io
import json
import logging
import re
from typing import Any

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

logger = logging.getLogger(__name__)

PNG_MAGIC = b"\x89PNG\r\n\x1a\n"
MAX_IMAGE_BYTES = 20 * 1024 * 1024
MAX_ELEMENTS = 60
MAX_TEXT = 2000
EMU_PER_PX = 9525  # 914400 EMU per inch / 96 DPI
_SAFE_KEY = re.compile(r"^[A-Za-z0-9._-]{1,80}$")

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def parse_layers_json(raw: str | None) -> dict[str, Any] | None:
    if not raw or not str(raw).strip():
        return None
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return None
    if not isinstance(data, dict):
        return None
    width = _int(data.get("width"), 1080, lo=40, hi=8000)
    height = _int(data.get("height"), 1080, lo=40, hi=8000)
    elements_in = data.get("elements")
    if not isinstance(elements_in, list):
        elements_in = []
    elements: list[dict[str, Any]] = []
    for item in elements_in[:MAX_ELEMENTS]:
        if isinstance(item, dict):
            elements.append(item)
    cover_key = data.get("cover_asset_key")
    if not isinstance(cover_key, str) or not _SAFE_KEY.match(cover_key.strip()):
        cover_key = None
    else:
        cover_key = cover_key.strip()
    return {
        "width": width,
        "height": height,
        "brand_logo": bool(data.get("brand_logo")),
        "cover_asset_key": cover_key,
        "elements": elements,
    }


def lookup_image(images: dict[str, bytes], key: str | None) -> bytes | None:
    if not key:
        return None
    direct = images.get(key)
    if direct:
        return direct
    stem = key.rsplit(".", 1)[0]
    return images.get(stem)


def build_pptx_from_layers(layers: dict[str, Any], images: dict[str, bytes]) -> bytes | None:
    """Return PPTX bytes with separate pictures + text frames, or None."""
    try:
        return _build_pptx(layers, images)
    except Exception:
        logger.warning("Canva PPTX layer build failed", exc_info=True)
        return None


def _build_pptx(layers: dict[str, Any], images: dict[str, bytes]) -> bytes | None:
    width = _int(layers.get("width"), 1080, lo=40, hi=8000)
    height = _int(layers.get("height"), 1080, lo=40, hi=8000)
    prs = Presentation()
    prs.slide_width = Emu(width * EMU_PER_PX)
    prs.slide_height = Emu(height * EMU_PER_PX)
    slide = prs.slides.add_slide(_blank_layout(prs))

    added = 0
    cover_key = layers.get("cover_asset_key")
    cover = lookup_image(images, cover_key if isinstance(cover_key, str) else None)
    if cover:
        fitted = _cover_fit_png(cover, width, height)
        if fitted and _add_picture(slide, fitted, 0, 0, width, height):
            added += 1

    if layers.get("brand_logo"):
        badge = max(28, round(width * 0.06))
        bx = round(width * 0.06)
        by = round(height * 0.06)
        bw = round(badge * 1.6)
        bh = badge
        if _add_round_rect(slide, bx, by, bw, bh, "#ffffff"):
            added += 1
        if _add_textbox(
            slide,
            "IH",
            x=bx,
            y=by,
            w=bw,
            h=bh,
            font_px=max(12, round(badge * 0.45)),
            bold=True,
            align="center",
            color="#111827",
            valign="middle",
        ):
            added += 1

    raw_elements = layers.get("elements")
    elements = raw_elements if isinstance(raw_elements, list) else []
    for item in sorted(elements, key=_z_key)[:MAX_ELEMENTS]:
        if not isinstance(item, dict):
            continue
        kind = str(item.get("type") or "").upper()
        if kind == "TEXT":
            if _add_text_element(slide, item):
                added += 1
        elif kind == "IMAGE":
            if _add_image_element(slide, item, images):
                added += 1
        elif kind == "BUTTON":
            if _add_button_element(slide, item):
                added += 1
        elif kind == "METRIC_GROUP":
            added += _add_metric_group(slide, item)

    if added < 1:
        return None
    buffer = io.BytesIO()
    prs.save(buffer)
    payload = buffer.getvalue()
    return payload if payload[:2] == b"PK" else None


def _blank_layout(prs: Presentation) -> Any:
    layouts = list(prs.slide_layouts)
    for layout in layouts:
        if "blank" in str(getattr(layout, "name", "")).lower():
            return layout
    return layouts[6] if len(layouts) > 6 else layouts[-1]


def _z_key(item: dict[str, Any]) -> tuple[int, str]:
    return (_int(item.get("zIndex"), 0, lo=-10_000, hi=10_000), str(item.get("id") or ""))


def _int(value: Any, default: int, *, lo: int, hi: int) -> int:
    try:
        number = int(round(float(value)))
    except (TypeError, ValueError):
        number = default
    return max(lo, min(hi, number))


def _emu(px: int) -> Emu:
    return Emu(max(1, int(px)) * EMU_PER_PX)


def _rgb(color: Any, fallback: str = "ffffff") -> RGBColor:
    raw = str(color or "").strip().lstrip("#")
    if len(raw) == 3 and all(c in "0123456789abcdefABCDEF" for c in raw):
        raw = "".join(c * 2 for c in raw)
    if len(raw) != 6 or any(c not in "0123456789abcdefABCDEF" for c in raw):
        raw = fallback
    return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))


def _clip_text(value: Any) -> str:
    return str(value or "").replace("\x00", "")[:MAX_TEXT]


def _px_to_pt(font_px: int) -> Pt:
    return Pt(max(8, min(200, round(font_px * 0.75))))


def _align(value: Any) -> PP_ALIGN:
    token = str(value or "left").lower()
    if token == "center":
        return PP_ALIGN.CENTER
    if token == "right":
        return PP_ALIGN.RIGHT
    return PP_ALIGN.LEFT


def _image_to_png(data: bytes) -> bytes | None:
    if not data or len(data) > MAX_IMAGE_BYTES:
        return None
    if data[:8] == PNG_MAGIC:
        return data
    try:
        img = Image.open(io.BytesIO(data))
        img = img.convert("RGBA")
        out = io.BytesIO()
        img.save(out, format="PNG")
        png = out.getvalue()
    except Exception:
        return None
    if not png or len(png) > MAX_IMAGE_BYTES:
        return None
    return png


def _cover_fit_png(data: bytes, box_w: int, box_h: int) -> bytes | None:
    png = _image_to_png(data)
    if not png:
        return None
    box_w = max(1, int(box_w))
    box_h = max(1, int(box_h))
    try:
        img = Image.open(io.BytesIO(png)).convert("RGBA")
    except Exception:
        return None
    if img.width < 1 or img.height < 1:
        return None
    scale = max(box_w / img.width, box_h / img.height)
    new_w = max(1, int(round(img.width * scale)))
    new_h = max(1, int(round(img.height * scale)))
    resized = img.resize((new_w, new_h), Image.Resampling.LANCZOS)
    left = max(0, (new_w - box_w) // 2)
    top = max(0, (new_h - box_h) // 2)
    cropped = resized.crop((left, top, left + box_w, top + box_h))
    out = io.BytesIO()
    cropped.save(out, format="PNG")
    return out.getvalue()


def _add_picture(slide: Any, png: bytes, x: int, y: int, w: int, h: int) -> bool:
    if w < 1 or h < 1:
        return False
    stream = io.BytesIO(png)
    stream.name = "layer.png"
    slide.shapes.add_picture(stream, _emu(x), _emu(y), _emu(w), _emu(h))
    return True


def _add_round_rect(slide: Any, x: int, y: int, w: int, h: int, fill: str) -> bool:
    if w < 8 or h < 8:
        return False
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, _emu(x), _emu(y), _emu(w), _emu(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill, "ffffff")
    shape.line.fill.background()
    with contextlib.suppress(Exception):
        shape.adjustments[0] = 0.5
    return True


def _add_textbox(
    slide: Any,
    text: str,
    *,
    x: int,
    y: int,
    w: int,
    h: int,
    font_px: int,
    bold: bool,
    align: str,
    color: str,
    valign: str = "top",
) -> bool:
    content = _clip_text(text).strip()
    if not content or w < 8 or h < 8:
        return False
    box = slide.shapes.add_textbox(_emu(x), _emu(y), _emu(w), _emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    with contextlib.suppress(Exception):
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE if valign == "middle" else MSO_ANCHOR.TOP
    paragraph = tf.paragraphs[0]
    paragraph.text = content
    paragraph.alignment = _align(align)
    paragraph.font.size = _px_to_pt(font_px)
    paragraph.font.bold = bool(bold)
    paragraph.font.color.rgb = _rgb(color)
    paragraph.font.name = "Arial"
    return True


def _add_text_element(slide: Any, item: dict[str, Any]) -> bool:
    return _add_textbox(
        slide,
        item.get("content"),
        x=_int(item.get("x"), 0, lo=0, hi=8000),
        y=_int(item.get("y"), 0, lo=0, hi=8000),
        w=_int(item.get("width"), 100, lo=8, hi=8000),
        h=_int(item.get("height"), 40, lo=8, hi=8000),
        font_px=_int(item.get("fontSize"), 24, lo=8, hi=400),
        bold=str(item.get("fontWeight") or "").lower() == "bold",
        align=str(item.get("align") or "left"),
        color=str(item.get("color") or "#ffffff"),
    )


def _add_image_element(slide: Any, item: dict[str, Any], images: dict[str, bytes]) -> bool:
    key = item.get("asset_key")
    raw = lookup_image(images, key if isinstance(key, str) else None)
    if not raw:
        return False
    x = _int(item.get("x"), 0, lo=0, hi=8000)
    y = _int(item.get("y"), 0, lo=0, hi=8000)
    w = _int(item.get("width"), 100, lo=8, hi=8000)
    h = _int(item.get("height"), 100, lo=8, hi=8000)
    fitted = _cover_fit_png(raw, w, h)
    if not fitted:
        return False
    return _add_picture(slide, fitted, x, y, w, h)


def _add_button_element(slide: Any, item: dict[str, Any]) -> bool:
    x = _int(item.get("x"), 0, lo=0, hi=8000)
    y = _int(item.get("y"), 0, lo=0, hi=8000)
    w = _int(item.get("width"), 160, lo=8, hi=8000)
    h = _int(item.get("height"), 40, lo=8, hi=8000)
    added = _add_round_rect(slide, x, y, w, h, str(item.get("backgroundColor") or "#ffffff"))
    label = _clip_text(item.get("label")).strip()
    if label:
        added = (
            _add_textbox(
                slide,
                label,
                x=x,
                y=y,
                w=w,
                h=h,
                font_px=max(12, round(h * 0.42)),
                bold=True,
                align="center",
                color=str(item.get("textColor") or "#111827"),
                valign="middle",
            )
            or added
        )
    return added


def _add_metric_group(slide: Any, item: dict[str, Any]) -> int:
    metrics = item.get("metrics")
    if not isinstance(metrics, list) or not metrics:
        return 0
    x = _int(item.get("x"), 0, lo=0, hi=8000)
    y = _int(item.get("y"), 0, lo=0, hi=8000)
    w = _int(item.get("width"), 200, lo=8, hi=8000)
    h = _int(item.get("height"), 80, lo=8, hi=8000)
    color = str(item.get("color") or "#ffffff")
    layout = str(item.get("layout") or "horizontal")
    count = max(1, min(8, len(metrics)))
    added = 0
    for index, metric in enumerate(metrics[:count]):
        if not isinstance(metric, dict):
            continue
        if layout == "stacked":
            row_h = max(8, h // count)
            mx, my, mw, mh = x, y + index * row_h, w, row_h
        else:
            col_w = max(8, w // count)
            mx, my, mw, mh = x + index * col_w, y, col_w, h
        value_h = max(16, round(mh * 0.55))
        if _add_textbox(
            slide,
            metric.get("display_value"),
            x=mx,
            y=my,
            w=mw,
            h=value_h,
            font_px=max(14, round(value_h * 0.55)),
            bold=True,
            align="left",
            color=color,
        ):
            added += 1
        if _add_textbox(
            slide,
            metric.get("label"),
            x=mx,
            y=my + value_h,
            w=mw,
            h=max(12, mh - value_h),
            font_px=max(10, round((mh - value_h) * 0.45)),
            bold=False,
            align="left",
            color=color,
        ):
            added += 1
    return added
