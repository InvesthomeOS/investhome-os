"""REIT membership: workbook first, then the required 10-participant set.

Source Excel still has 12 REIT rows. Two incorrect REIT relationships are
removed from CRM only. Contacts are never deleted. Amounts are never invented.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from investhome_api.services.crm.bitrix_import import BitrixBundle, BitrixSourceRow
from investhome_api.services.crm.bitrix_project_aliases import (
    BitrixProjectGroup,
    fold_alias,
    resolve_project_group,
)

EXPECTED_REIT_PARTICIPANTS = 10
EXPECTED_AGREEMENT_TOTAL = 65
SOURCE_AGREEMENT_ROWS = 67
EXPECTED_AGREEMENT_COUNTS: dict[str, int] = {
    BitrixProjectGroup.KST_1307.value: 6,
    BitrixProjectGroup.PENN_1313.value: 2,
    BitrixProjectGroup.HPL_1812.value: 17,
    BitrixProjectGroup.ONTARIO_2319.value: 9,
    BitrixProjectGroup.REIT.value: EXPECTED_REIT_PARTICIPANTS,
    BitrixProjectGroup.THE_TEMPLE.value: 1,
    BitrixProjectGroup.UNILOFT.value: 20,
}


def _source_id(row: BitrixSourceRow) -> str | None:
    bitrix_id = (row.bitrix_id or "").strip()
    if not bitrix_id:
        return None
    return f"{row.source_file}:{bitrix_id}"


def _fold_external_id(value: str | None) -> str:
    text = (value or "").strip()
    if not text:
        return ""
    file_part, sep, bitrix_id = text.rpartition(":")
    if not sep:
        return fold_alias(text)
    return f"{fold_alias(file_part)}:{bitrix_id.strip()}"


def reit_source_rows(bundle: BitrixBundle) -> list[BitrixSourceRow]:
    rows: list[BitrixSourceRow] = []
    for row in bundle.rows:
        if row.role != "agreements":
            continue
        if resolve_project_group(row.source_file) != BitrixProjectGroup.REIT:
            continue
        rows.append(row)
    return rows


def is_excluded_reit_source_id(source_external_id: str | None, excluded: set[str]) -> bool:
    if not source_external_id or not excluded:
        return False
    needle = _fold_external_id(source_external_id)
    if not needle:
        return False
    folded_excluded = {_fold_external_id(item) for item in excluded}
    if needle in folded_excluded:
        return True
    bitrix_id = source_external_id.rsplit(":", 1)[-1].strip()
    if not bitrix_id:
        return False
    file_part = source_external_id.rsplit(":", 1)[0]
    if resolve_project_group(file_part) != BitrixProjectGroup.REIT:
        return False
    return any(item.rsplit(":", 1)[-1].strip() == bitrix_id for item in excluded)


def _cell_text(value: object) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    text = str(value).strip()
    return text or None


def _workbook_reit_keep_ids(path: Path | None, source_ids: set[str]) -> set[str] | None:
    if path is None or not path.exists():
        return None
    workbook = load_workbook(path, data_only=True, read_only=True)
    keep: set[str] = set()
    found_reit_sheet = False
    try:
        for sheet in workbook.worksheets:
            title = fold_alias(sheet.title)
            if "missing units" not in title and "eksik daire" not in title:
                continue
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = [fold_alias(str(item) if item is not None else "") for item in next(rows)]
            except StopIteration:
                continue
            project_idx = next(
                (i for i, key in enumerate(headers) if key in {"agreement project", "project", "proje"}),
                None,
            )
            file_idx = next(
                (
                    i
                    for i, key in enumerate(headers)
                    if key in {"source excel filename", "source file", "filename"}
                ),
                None,
            )
            id_idx = next(
                (
                    i
                    for i, key in enumerate(headers)
                    if key in {"bitrix external source id", "bitrix id", "source id", "external id"}
                ),
                None,
            )
            class_idx = next((i for i, key in enumerate(headers) if key in {"class", "sinif"}), None)
            for raw in rows:
                values = list(raw)
                project = _cell_text(values[project_idx]) if project_idx is not None and project_idx < len(values) else None
                source_file = _cell_text(values[file_idx]) if file_idx is not None and file_idx < len(values) else None
                klass = _cell_text(values[class_idx]) if class_idx is not None and class_idx < len(values) else None
                group = resolve_project_group(project) or resolve_project_group(source_file)
                if group != BitrixProjectGroup.REIT:
                    continue
                if klass and klass.strip().upper() not in {"A", "CLASS A"}:
                    continue
                found_reit_sheet = True
                external = _cell_text(values[id_idx]) if id_idx is not None and id_idx < len(values) else None
                if not external:
                    continue
                if external in source_ids:
                    keep.add(external)
                    continue
                for source_id in source_ids:
                    if is_excluded_reit_source_id(external, {source_id}):
                        keep.add(source_id)
                        break
    finally:
        workbook.close()
    if not found_reit_sheet:
        return None
    return keep


def detect_excluded_reit_source_ids(
    bundle: BitrixBundle,
    *,
    rescue_xlsx: Path | None = None,
) -> tuple[set[str], str]:
    rows = reit_source_rows(bundle)
    source_by_id = {key: row for row in rows if (key := _source_id(row))}
    source_ids = set(source_by_id)
    if not source_ids:
        return set(), "no_reit_source_rows"

    workbook_keep = _workbook_reit_keep_ids(rescue_xlsx, source_ids)
    if workbook_keep is not None:
        removed = source_ids - workbook_keep
        if (
            len(workbook_keep) == EXPECTED_REIT_PARTICIPANTS
            and len(removed) == len(source_ids) - EXPECTED_REIT_PARTICIPANTS
        ):
            return removed, "workbook_membership"

    extra = len(source_ids) - EXPECTED_REIT_PARTICIPANTS
    if extra <= 0:
        return set(), "already_at_or_below_expected"
    blank = {
        source_id
        for source_id, row in source_by_id.items()
        if not (row.payment_amount or "").strip()
    }
    if len(blank) == extra:
        return blank, "operator_count_blank_gelir"
    return set(), "unresolved"
