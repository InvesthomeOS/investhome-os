"""Phase 6E: manual rescue backfill, REIT Gelir, junk reasons, project reconciliation.

Does not create contacts or agreements. Does not invent phones, emails, or units.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session, attributes

from investhome_api.models.crm_agreement import CrmAgreement
from investhome_api.models.crm_contact import CrmContact, CrmContactStatus
from investhome_api.services.crm.bitrix_import import BITRIX_SOURCE, BitrixBundle, BitrixSourceRow
from investhome_api.services.crm.bitrix_project_aliases import (
    BITRIX_PROJECT_GROUP_LABELS,
    BitrixProjectGroup,
    fold_alias,
    resolve_project_group,
)
from investhome_api.services.crm.bitrix_reit_membership import (
    EXPECTED_AGREEMENT_COUNTS,
    EXPECTED_AGREEMENT_TOTAL,
    SOURCE_AGREEMENT_ROWS,
    detect_excluded_reit_source_ids,
    is_excluded_reit_source_id,
)
from investhome_api.services.crm.identity import (
    displayable_phone,
    is_zero_masked_phone,
    normalize_full_name,
    normalize_valid_email,
    parse_phone,
)

_METADATA_KEY = "bitrix_import"
_INTENTIONALLY_BLANK = {
    "blank",
    "yok",
    "missing",
    "n/a",
    "na",
    "none",
    "null",
    "intentionally blank",
    "confirmed missing",
    "phone missing",
    "leave blank",
    "-",
    "—",
    "yoktur",
}
_CORRECTED_PHONE = {
    "corrected phone",
    "manual phone",
    "yeni telefon",
    "phone corrected",
    "corrected_phone",
    "telefon duzeltme",
    "telefon duzeltilmis",
}
_CORRECTED_EMAIL = {
    "corrected email",
    "manual email",
    "yeni eposta",
    "yeni e posta",
    "email corrected",
    "corrected_email",
}
_CORRECTED_NAME = {
    "corrected name",
    "manual name",
    "yeni ad",
    "corrected full name",
    "corrected_name",
}
_NOTE_HEADERS = {
    "notes",
    "note",
    "comment",
    "comments",
    "manual notes",
    "manual comments",
    "bitrix comments",
    "not",
    "notlar",
    "yorum",
    "yorumlar",
}
_VERIFIED_HEADERS = {
    "verified",
    "manual verification",
    "reviewed",
    "dogrulandi",
    "manuel dogrulama",
}
_PHONE_MISSING_HEADERS = {
    "phone confirmed missing",
    "phone missing",
    "telefon yok",
    "telefon eksik",
}


@dataclass
class FinalNormalizationReport:
    dry_run: bool
    contacts_before: int
    contacts_after: int
    agreements_before: int
    agreements_after: int
    project_groups_before: dict[str, int] = field(default_factory=dict)
    project_groups_after: dict[str, int] = field(default_factory=dict)
    rescue_corrections_applied: int = 0
    rescue_notes_preserved: int = 0
    rescue_blank_skipped: int = 0
    rescue_verified_flagged: int = 0
    reit_investment_updated: int = 0
    reit_missing_gelir: list[dict[str, Any]] = field(default_factory=list)
    reit_membership_reason: str = ""
    reit_excluded_source_ids: list[str] = field(default_factory=list)
    reit_membership_removed: list[dict[str, Any]] = field(default_factory=list)
    reit_contacts_preserved: int = 0
    junk_source_rows: int = 0
    junk_source_with_reason: int = 0
    junk_reasons_updated: int = 0
    junk_reasons_missing: int = 0
    wrong_project_found: int = 0
    wrong_project_repaired: int = 0
    wrong_contact_found: int = 0
    wrong_contact_repaired: int = 0
    wrong_contact_unresolved: list[dict[str, Any]] = field(default_factory=list)
    legitimate_multi_project: list[dict[str, Any]] = field(default_factory=list)
    residual_multi_project: list[dict[str, Any]] = field(default_factory=list)
    masked_phones_cleared: int = 0
    blank_phone_agreement_contacts: list[dict[str, Any]] = field(default_factory=list)
    reconciliation_path: str | None = None
    notes: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    def format_console(self) -> str:
        payload = self.to_dict()
        lines = ["BITRIX FINAL NORMALIZATION", ""]
        for key, value in payload.items():
            if key in {"blank_phone_agreement_contacts", "legitimate_multi_project", "residual_multi_project"}:
                lines.append(f"  {key}: {len(value) if isinstance(value, list) else value}")
            elif key.endswith("_samples"):
                continue
            else:
                lines.append(f"  {key}: {value}")
        return "\n".join(lines) + "\n"


def _clip(value: str | None, size: int) -> str | None:
    text = (value or "").strip()
    return text[:size] if text else None


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _project_counts(rows: list[CrmAgreement]) -> dict[str, int]:
    counts: Counter[str] = Counter()
    for row in rows:
        counts[row.project_group] += 1
    return dict(sorted(counts.items()))


def _source_id(row: BitrixSourceRow) -> str | None:
    bitrix_id = (row.bitrix_id or "").strip()
    if not bitrix_id:
        return None
    return f"{row.source_file}:{bitrix_id}"


def _is_imported_agreement(row: CrmAgreement) -> bool:
    if (row.source or "").strip().casefold() == BITRIX_SOURCE:
        return True
    meta = row.metadata_json if isinstance(row.metadata_json, dict) else {}
    return bool(meta.get("imported_historical_agreement") or meta.get("source") == "Bitrix")


def _bitrix_meta(contact: CrmContact) -> dict[str, Any]:
    raw = contact.metadata_json if isinstance(contact.metadata_json, dict) else {}
    nested = raw.get(_METADATA_KEY)
    return dict(nested) if isinstance(nested, dict) else {}


def _set_bitrix_meta(contact: CrmContact, updates: dict[str, Any]) -> None:
    raw = dict(contact.metadata_json) if isinstance(contact.metadata_json, dict) else {}
    nested = dict(raw.get(_METADATA_KEY)) if isinstance(raw.get(_METADATA_KEY), dict) else {}
    nested.update(updates)
    raw[_METADATA_KEY] = nested
    contact.metadata_json = raw
    attributes.flag_modified(contact, "metadata_json")


def _agreement_meta(row: CrmAgreement) -> dict[str, Any]:
    return dict(row.metadata_json) if isinstance(row.metadata_json, dict) else {}


def _set_agreement_meta(row: CrmAgreement, metadata: dict[str, Any]) -> None:
    row.metadata_json = metadata
    attributes.flag_modified(row, "metadata_json")


def _truthy(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    text = str(value).strip().casefold()
    return text in {"1", "true", "yes", "y", "evet", "x", "ok", "verified", "dogrulandi"}


def _cell_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    text = str(value).strip()
    return text or None


def _header_map(headers: list[Any]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    for index, header in enumerate(headers):
        key = fold_alias(str(header) if header is not None else "")
        if key and key not in mapped:
            mapped[key] = index
    return mapped


def _column(mapped: dict[str, int], aliases: set[str], row: list[Any]) -> str | None:
    for alias in aliases:
        index = mapped.get(alias)
        if index is None or index >= len(row):
            continue
        text = _cell_text(row[index])
        if text:
            return text
    return None


def _contacts_by_email(contacts: list[CrmContact]) -> dict[str, list[CrmContact]]:
    index: dict[str, list[CrmContact]] = defaultdict(list)
    for contact in contacts:
        emails = [contact.primary_email, *(contact.secondary_emails or [])]
        seen: set[str] = set()
        for email in emails:
            key = normalize_valid_email(email)
            if not key or key in seen:
                continue
            seen.add(key)
            index[key].append(contact)
    return index


def _match_contact_for_row(
    row: BitrixSourceRow,
    *,
    by_email: dict[str, list[CrmContact]],
    by_id: dict[UUID, CrmContact],
) -> CrmContact | None:
    email = normalize_valid_email(row.email)
    if not email:
        return None
    hits = list(by_email.get(email) or [])
    if not hits:
        return None
    name_key = normalize_full_name(row.full_name)
    named = [item for item in hits if normalize_full_name(item.display_name) == name_key]
    if len(named) == 1:
        return named[0]
    if len(hits) == 1:
        return hits[0]
    return None


def _identity_mismatch(row: BitrixSourceRow, contact: CrmContact | None) -> bool:
    if contact is None:
        return False
    source_name = normalize_full_name(row.full_name)
    contact_name = normalize_full_name(contact.display_name)
    source_email = normalize_valid_email(row.email)
    contact_email = normalize_valid_email(contact.primary_email)
    if source_name and contact_name and source_name != contact_name:
        if source_email and contact_email and source_email != contact_email:
            return True
    return False


def _load_rescue_rows(path: Path) -> list[dict[str, Any]]:
    if not path.exists():
        return []
    workbook = load_workbook(path, data_only=True, read_only=True)
    records: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = list(next(rows))
            except StopIteration:
                continue
            mapped = _header_map(headers)
            for raw in rows:
                values = list(raw)
                if not any(_cell_text(item) for item in values):
                    continue
                records.append(
                    {
                        "sheet": sheet.title,
                        "mapped": mapped,
                        "values": values,
                    }
                )
    finally:
        workbook.close()
    return records


def _apply_rescue(
    db: Session,
    *,
    contacts: list[CrmContact],
    agreements: list[CrmAgreement],
    rescue_path: Path | None,
    dry_run: bool,
    report: FinalNormalizationReport,
) -> None:
    if rescue_path is None:
        report.notes.append("rescue_workbook_not_provided")
        return
    if not rescue_path.exists():
        report.notes.append(f"rescue_workbook_missing:{rescue_path}")
        return
    rows = _load_rescue_rows(rescue_path)
    report.notes.append(f"rescue_workbook_rows:{len(rows)}")
    by_email = _contacts_by_email(contacts)
    by_external: dict[str, CrmContact] = {}
    for contact in contacts:
        meta = _bitrix_meta(contact)
        for value in meta.get("external_ids") or []:
            key = str(value).strip()
            if key:
                by_external[key] = contact
    by_agreement_id = {
        (row.source_external_id or "").strip(): row
        for row in agreements
        if row.source_external_id
    }
    for item in rows:
        mapped: dict[str, int] = item["mapped"]
        values: list[Any] = item["values"]
        external = _column(
            mapped,
            {"bitrix external source id", "bitrix id", "source id", "external id"},
            values,
        )
        email = normalize_valid_email(
            _column(mapped, {"email", "e posta", "e-posta"}, values)
        )
        name = _column(mapped, {"full name", "source full name", "os contact name", "display name"}, values)
        contact: CrmContact | None = None
        if external and external in by_external:
            contact = by_external[external]
        if contact is None and external and external in by_agreement_id:
            agreement = by_agreement_id[external]
            contact = next((item for item in contacts if item.id == agreement.contact_id), None)
        if contact is None and email:
            hits = by_email.get(email) or []
            named = [item for item in hits if normalize_full_name(item.display_name) == normalize_full_name(name)]
            if len(named) == 1:
                contact = named[0]
            elif len(hits) == 1:
                contact = hits[0]
        if contact is None:
            continue

        verified = _truthy(_column(mapped, _VERIFIED_HEADERS, values))
        phone_missing = _truthy(_column(mapped, _PHONE_MISSING_HEADERS, values))
        corrected_phone = _column(mapped, _CORRECTED_PHONE, values)
        corrected_email = _column(mapped, _CORRECTED_EMAIL, values)
        corrected_name = _column(mapped, _CORRECTED_NAME, values)
        note = _column(mapped, _NOTE_HEADERS, values)

        changed = False
        if corrected_phone:
            token = fold_alias(corrected_phone)
            if token in _INTENTIONALLY_BLANK or phone_missing:
                if contact.primary_phone:
                    if not dry_run:
                        _set_bitrix_meta(
                            contact,
                            {
                                "raw_source_phone": contact.primary_phone,
                                "phone_display_suppressed": True,
                                "manual_bitrix_verification": True,
                            },
                        )
                        contact.primary_phone = None
                    changed = True
            elif displayable_phone(corrected_phone):
                parsed = parse_phone(corrected_phone)
                new_phone = parsed.e164 if parsed and parsed.e164 and not parsed.suspicious else corrected_phone[:50]
                if new_phone != contact.primary_phone:
                    if not dry_run:
                        contact.primary_phone = new_phone
                    changed = True
            else:
                report.rescue_blank_skipped += 1
        elif not _column(mapped, _CORRECTED_PHONE, values):
            pass

        if corrected_email:
            token = fold_alias(corrected_email)
            if token in _INTENTIONALLY_BLANK:
                report.rescue_blank_skipped += 1
            else:
                valid = normalize_valid_email(corrected_email)
                if valid and valid != normalize_valid_email(contact.primary_email):
                    if not dry_run:
                        contact.primary_email = valid
                    changed = True
        if corrected_name:
            token = fold_alias(corrected_name)
            if token in _INTENTIONALLY_BLANK:
                report.rescue_blank_skipped += 1
            elif corrected_name != contact.display_name:
                if not dry_run:
                    contact.display_name = corrected_name[:255]
                changed = True
        if note:
            existing = (contact.notes or "").strip()
            if note not in existing:
                if not dry_run:
                    contact.notes = f"{existing}\n{note}".strip() if existing else note
                report.rescue_notes_preserved += 1
                changed = True
        if verified or changed:
            if not dry_run:
                _set_bitrix_meta(contact, {"manual_bitrix_verification": True, "manual_verified_at": _now()})
            report.rescue_verified_flagged += 1
        if changed:
            report.rescue_corrections_applied += 1


def _apply_reit_membership(
    db: Session,
    *,
    agreements: list[CrmAgreement],
    contacts: list[CrmContact],
    excluded: set[str],
    dry_run: bool,
    report: FinalNormalizationReport,
) -> None:
    if not excluded:
        return
    contacts_by_id = {item.id: item for item in contacts}
    remaining: list[CrmAgreement] = []
    preserved: set[UUID] = set()
    for agreement in agreements:
        if agreement.project_group != BitrixProjectGroup.REIT.value or not is_excluded_reit_source_id(
            agreement.source_external_id, excluded
        ):
            remaining.append(agreement)
            continue
        contact = contacts_by_id.get(agreement.contact_id)
        report.reit_membership_removed.append(
            {
                "source_external_id": agreement.source_external_id,
                "contact_id": str(agreement.contact_id),
                "contact_name": contact.display_name if contact else None,
                "contact_email": contact.primary_email if contact else None,
            }
        )
        if contact is not None:
            preserved.add(contact.id)
            if not dry_run:
                _set_bitrix_meta(
                    contact,
                    {
                        "reit_membership_removed": True,
                        "reit_membership_removed_source_id": agreement.source_external_id,
                        "reit_membership_removed_at": _now(),
                    },
                )
        if not dry_run:
            db.delete(agreement)
    report.reit_contacts_preserved = len(preserved)
    if not dry_run:
        agreements[:] = remaining


def _apply_reit_investment(
    *,
    agreements: list[CrmAgreement],
    agreement_rows: list[BitrixSourceRow],
    excluded: set[str],
    dry_run: bool,
    report: FinalNormalizationReport,
) -> None:
    by_source = {(row.source, row.source_external_id): row for row in agreements if row.source_external_id}
    for row in agreement_rows:
        group = resolve_project_group(row.source_file)
        if group != BitrixProjectGroup.REIT:
            continue
        external = _source_id(row)
        if is_excluded_reit_source_id(external, excluded):
            continue
        existing = by_source.get((BITRIX_SOURCE, external)) if external else None
        if existing is None:
            report.notes.append(f"reit_source_unmatched:{row.source_file}:{row.bitrix_id}")
            continue
        amount = _clip(row.payment_amount, 80)
        if not amount:
            report.reit_missing_gelir.append(
                {
                    "source_external_id": existing.source_external_id,
                    "name": row.full_name,
                    "bitrix_id": row.bitrix_id,
                }
            )
        if dry_run:
            if amount and existing.investment_amount != amount:
                report.reit_investment_updated += 1
            continue
        if amount and existing.investment_amount != amount:
            existing.investment_amount = amount
            report.reit_investment_updated += 1
        if existing.unit_number:
            existing.unit_number = None
        meta = _agreement_meta(existing)
        fields = dict(meta.get("agreement_fields") or {}) if isinstance(meta.get("agreement_fields"), dict) else {}
        meta_changed = False
        if amount and meta.get("investment_amount") != amount:
            meta["investment_amount"] = amount
            fields["payment_amount"] = amount
            meta_changed = True
        if fields.pop("unit_number", None) is not None:
            meta_changed = True
        if meta.pop("unit_number", None) is not None:
            meta_changed = True
        if not meta.get("reit_investment"):
            meta["reit_investment"] = True
            meta_changed = True
        if meta_changed:
            meta["agreement_fields"] = {key: value for key, value in fields.items() if value}
            _set_agreement_meta(existing, meta)


def _apply_junk_reasons(
    *,
    contacts: list[CrmContact],
    junk_rows: list[BitrixSourceRow],
    dry_run: bool,
    report: FinalNormalizationReport,
) -> None:
    report.junk_source_rows = len(junk_rows)
    report.junk_source_with_reason = sum(1 for row in junk_rows if (row.junk_reason or "").strip())
    by_email = _contacts_by_email(contacts)
    by_external: dict[str, list[CrmContact]] = defaultdict(list)
    for contact in contacts:
        meta = _bitrix_meta(contact)
        for value in meta.get("external_ids") or []:
            key = str(value).strip()
            if key:
                by_external[key].append(contact)
        for source_file in meta.get("source_files") or []:
            for value in meta.get("external_ids") or []:
                compound = f"{source_file}:{value}"
                by_external[compound].append(contact)

    matched: set[UUID] = set()
    updated: set[UUID] = set()
    for row in junk_rows:
        reason = _clip(row.junk_reason, 255)
        if not reason:
            continue
        candidates: list[CrmContact] = []
        external = _source_id(row)
        if external:
            candidates.extend(by_external.get(external, []))
        if row.bitrix_id:
            candidates.extend(by_external.get(row.bitrix_id.strip(), []))
        email = normalize_valid_email(row.email)
        if email:
            candidates.extend(by_email.get(email, []))
        unique: dict[UUID, CrmContact] = {}
        for contact in candidates:
            unique[contact.id] = contact
        if len(unique) != 1:
            continue
        contact = next(iter(unique.values()))
        matched.add(contact.id)
        historical = _bitrix_meta(contact)
        reasons = [str(item) for item in (historical.get("historical_junk_reasons") or []) if item]
        added_historical = False
        if reason not in reasons:
            reasons.append(reason)
            added_historical = True
        fill_first_class = not (contact.junk_reason or "").strip()
        if not fill_first_class and not added_historical and historical.get("junk_reason"):
            continue
        if not dry_run:
            if fill_first_class:
                contact.junk_reason = reason
            _set_bitrix_meta(
                contact,
                {
                    "junk_reason": contact.junk_reason or reason,
                    "historical_junk_reasons": reasons,
                    "historical_junk": True,
                },
            )
        if fill_first_class and contact.id not in updated:
            report.junk_reasons_updated += 1
            updated.add(contact.id)

    archived = [item for item in contacts if item.status == CrmContactStatus.ARCHIVED]
    report.junk_reasons_missing = sum(1 for item in archived if not (item.junk_reason or "").strip())


def _relink_wrong_contacts(
    *,
    agreements: list[CrmAgreement],
    agreement_rows: list[BitrixSourceRow],
    contacts: list[CrmContact],
    excluded: set[str],
    dry_run: bool,
    report: FinalNormalizationReport,
) -> None:
    by_source = {(row.source, row.source_external_id): row for row in agreements if row.source_external_id}
    by_id = {item.id: item for item in contacts}
    by_email = _contacts_by_email(contacts)
    for row in agreement_rows:
        external = _source_id(row)
        if is_excluded_reit_source_id(external, excluded):
            continue
        existing = by_source.get((BITRIX_SOURCE, external)) if external else None
        if existing is None or not _is_imported_agreement(existing):
            continue
        contact = by_id.get(existing.contact_id)
        if not _identity_mismatch(row, contact):
            continue
        report.wrong_contact_found += 1
        target = _match_contact_for_row(row, by_email=by_email, by_id=by_id)
        payload = {
            "source_external_id": existing.source_external_id,
            "source_name": row.full_name,
            "source_email": row.email,
            "current_contact": contact.display_name if contact else None,
            "current_email": contact.primary_email if contact else None,
            "expected_project": resolve_project_group(row.source_file).value
            if resolve_project_group(row.source_file)
            else None,
        }
        if target is None or target.id == existing.contact_id:
            report.wrong_contact_unresolved.append({**payload, "reason": "no_unique_existing_contact"})
            continue
        if not dry_run:
            existing.contact_id = target.id
            meta = _agreement_meta(existing)
            meta["contact_relinked_from"] = str(contact.id) if contact else None
            meta["contact_relinked_to"] = str(target.id)
            meta["contact_relink_reason"] = "source_email_name_mismatch"
            _set_agreement_meta(existing, meta)
        report.wrong_contact_repaired += 1


def _repair_wrong_projects(
    *,
    agreements: list[CrmAgreement],
    agreement_rows: list[BitrixSourceRow],
    excluded: set[str],
    dry_run: bool,
    report: FinalNormalizationReport,
) -> None:
    by_source = {(row.source, row.source_external_id): row for row in agreements if row.source_external_id}
    for row in agreement_rows:
        expected = resolve_project_group(row.source_file)
        if expected is None:
            continue
        external = _source_id(row)
        if is_excluded_reit_source_id(external, excluded):
            continue
        existing = by_source.get((BITRIX_SOURCE, external)) if external else None
        if existing is None or not _is_imported_agreement(existing):
            continue
        if existing.project_group == expected.value:
            continue
        report.wrong_project_found += 1
        if not dry_run:
            existing.project_group = expected.value
            meta = _agreement_meta(existing)
            meta["project_group"] = expected.value
            meta["project_label"] = BITRIX_PROJECT_GROUP_LABELS[expected]
            meta["project_from_source_file"] = True
            _set_agreement_meta(existing, meta)
        report.wrong_project_repaired += 1


def _clear_masked_phones(
    *,
    contacts: list[CrmContact],
    agreement_contact_ids: set[UUID],
    dry_run: bool,
    report: FinalNormalizationReport,
) -> None:
    seen: set[UUID] = set()
    for contact in contacts:
        if contact.id not in agreement_contact_ids:
            continue
        if not is_zero_masked_phone(contact.primary_phone):
            continue
        if contact.id in seen:
            continue
        seen.add(contact.id)
        fallback = displayable_phone(*(contact.secondary_phones or []))
        if not dry_run:
            _set_bitrix_meta(
                contact,
                {
                    "raw_source_phone": contact.primary_phone,
                    "phone_display_suppressed": True,
                    "phone_suppressed_reason": "zero_masked_excel",
                },
            )
            contact.primary_phone = fallback
            if contact.secondary_phones:
                contact.secondary_phones = [
                    item for item in contact.secondary_phones if displayable_phone(item)
                ] or None
        report.masked_phones_cleared += 1


def _classify_multi_project(
    *,
    agreements: list[CrmAgreement],
    contacts_by_id: dict[UUID, CrmContact],
    agreement_rows_by_id: dict[str, BitrixSourceRow],
    report: FinalNormalizationReport,
) -> None:
    by_contact: dict[UUID, list[CrmAgreement]] = defaultdict(list)
    for row in agreements:
        by_contact[row.contact_id].append(row)
    for contact_id, rows in by_contact.items():
        groups = sorted({item.project_group for item in rows})
        if len(groups) < 2:
            continue
        contact = contacts_by_id.get(contact_id)
        files = []
        legitimate = True
        for item in rows:
            source_row = agreement_rows_by_id.get(item.source_external_id or "")
            expected = resolve_project_group(
                str((_agreement_meta(item).get("source_file") or (source_row.source_file if source_row else "")))
            )
            if expected is None or expected.value != item.project_group:
                legitimate = False
            if source_row and _identity_mismatch(source_row, contact):
                legitimate = False
            files.append(item.source_external_id)
        payload = {
            "contact_id": str(contact_id),
            "name": contact.display_name if contact else None,
            "projects": groups,
            "source_ids": files,
        }
        if legitimate:
            report.legitimate_multi_project.append(payload)
        else:
            report.residual_multi_project.append(payload)


def _write_reconciliation(
    path: Path,
    *,
    agreements: list[CrmAgreement],
    contacts_by_id: dict[UUID, CrmContact],
    agreement_rows_by_id: dict[str, BitrixSourceRow],
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Project reconciliation"
    sheet.append(
        [
            "Canonical Contact ID",
            "Customer Name",
            "Source Excel",
            "Source Row / Bitrix ID",
            "Expected Project",
            "Current Project",
            "Match / Wrong Project",
            "Unit Number",
            "Investment Amount",
            "Phone Present",
            "Email Present",
            "Manual Verification Applied",
            "Notes",
        ]
    )
    for row in sorted(agreements, key=lambda item: (item.project_group, item.source_external_id or "")):
        contact = contacts_by_id.get(row.contact_id)
        meta = _agreement_meta(row)
        source_row = agreement_rows_by_id.get(row.source_external_id or "")
        source_file = str(meta.get("source_file") or (source_row.source_file if source_row else "") or "")
        expected = resolve_project_group(source_file)
        expected_label = BITRIX_PROJECT_GROUP_LABELS[expected] if expected else ""
        current_label = BITRIX_PROJECT_GROUP_LABELS.get(
            BitrixProjectGroup(row.project_group), row.project_group
        ) if row.project_group else ""
        match = "MATCH"
        notes: list[str] = []
        if expected and expected.value != row.project_group:
            match = "WRONG_PROJECT"
            notes.append("project_group_mismatch")
        if source_row and _identity_mismatch(source_row, contact):
            notes.append("wrong_contact")
        customer = (
            (source_row.full_name if source_row else None)
            or meta.get("source_name")
            or (contact.display_name if contact else None)
        )
        phone = displayable_phone(
            contact.primary_phone if contact else None,
            str(meta.get("contact_phone") or "") or None,
            source_row.phone if source_row else None,
        )
        email = (
            (source_row.email if source_row else None)
            or meta.get("contact_email")
            or (contact.primary_email if contact else None)
        )
        verified = bool(_bitrix_meta(contact).get("manual_bitrix_verification")) if contact else False
        is_reit = row.project_group == BitrixProjectGroup.REIT.value
        sheet.append(
            [
                str(row.contact_id),
                customer,
                source_file,
                row.source_external_id,
                expected_label,
                current_label,
                match,
                "" if is_reit else (row.unit_number or ""),
                row.investment_amount or (source_row.payment_amount if source_row else "") or "",
                "yes" if phone else "no",
                "yes" if email else "no",
                "yes" if verified else "no",
                "; ".join(notes),
            ]
        )
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def run_bitrix_final_normalization(
    db: Session,
    bundle: BitrixBundle,
    *,
    rescue_xlsx: Path | None = None,
    reconciliation_xlsx: Path | None = None,
    dry_run: bool = True,
) -> FinalNormalizationReport:
    contacts = list(db.query(CrmContact).all())
    agreements = list(db.query(CrmAgreement).all())
    report = FinalNormalizationReport(
        dry_run=dry_run,
        contacts_before=len(contacts),
        contacts_after=len(contacts),
        agreements_before=len(agreements),
        agreements_after=len(agreements),
        project_groups_before=_project_counts(agreements),
    )
    agreement_rows = [row for row in bundle.rows if row.role == "agreements"]
    junk_rows = [row for row in bundle.rows if row.role == "junk"]
    if len(agreement_rows) != SOURCE_AGREEMENT_ROWS:
        report.notes.append(f"source_agreement_rows:{len(agreement_rows)}")
    excluded_reit, membership_reason = detect_excluded_reit_source_ids(
        bundle, rescue_xlsx=rescue_xlsx
    )
    report.reit_membership_reason = membership_reason
    report.reit_excluded_source_ids = sorted(excluded_reit)

    _apply_rescue(
        db,
        contacts=contacts,
        agreements=agreements,
        rescue_path=rescue_xlsx,
        dry_run=dry_run,
        report=report,
    )
    _apply_reit_membership(
        db,
        agreements=agreements,
        contacts=contacts,
        excluded=excluded_reit,
        dry_run=dry_run,
        report=report,
    )
    _apply_reit_investment(
        agreements=agreements,
        agreement_rows=agreement_rows,
        excluded=excluded_reit,
        dry_run=dry_run,
        report=report,
    )
    _apply_junk_reasons(
        contacts=contacts,
        junk_rows=junk_rows,
        dry_run=dry_run,
        report=report,
    )
    _repair_wrong_projects(
        agreements=agreements,
        agreement_rows=agreement_rows,
        excluded=excluded_reit,
        dry_run=dry_run,
        report=report,
    )
    _relink_wrong_contacts(
        agreements=agreements,
        agreement_rows=agreement_rows,
        contacts=contacts,
        excluded=excluded_reit,
        dry_run=dry_run,
        report=report,
    )
    agreement_contact_ids = {row.contact_id for row in agreements}
    _clear_masked_phones(
        contacts=contacts,
        agreement_contact_ids=agreement_contact_ids,
        dry_run=dry_run,
        report=report,
    )

    if not dry_run:
        db.flush()
        contacts = list(db.query(CrmContact).all())
        agreements = list(db.query(CrmAgreement).all())

    contacts_by_id = {item.id: item for item in contacts}
    agreement_rows_by_id = {
        key: row for row in agreement_rows if (key := _source_id(row))
    }
    _classify_multi_project(
        agreements=agreements,
        contacts_by_id=contacts_by_id,
        agreement_rows_by_id=agreement_rows_by_id,
        report=report,
    )

    blank: dict[UUID, dict[str, Any]] = {}
    for row in agreements:
        contact = contacts_by_id.get(row.contact_id)
        if contact is None:
            continue
        if displayable_phone(contact.primary_phone):
            continue
        blank[contact.id] = {
            "contact_id": str(contact.id),
            "name": contact.display_name,
            "email": contact.primary_email,
        }
    report.blank_phone_agreement_contacts = list(blank.values())
    report.contacts_after = len(contacts)
    report.agreements_after = len(agreements)
    report.project_groups_after = _project_counts(agreements)
    if report.contacts_after != report.contacts_before:
        report.notes.append("contact_count_changed")
    if report.agreements_after != EXPECTED_AGREEMENT_TOTAL:
        report.notes.append(f"agreement_count:{report.agreements_after}")
    for group, expected in EXPECTED_AGREEMENT_COUNTS.items():
        actual = report.project_groups_after.get(group, 0)
        if actual != expected:
            report.notes.append(f"count_mismatch:{group}:{actual}:{expected}")

    if reconciliation_xlsx is not None:
        _write_reconciliation(
            reconciliation_xlsx,
            agreements=agreements,
            contacts_by_id=contacts_by_id,
            agreement_rows_by_id=agreement_rows_by_id,
        )
        report.reconciliation_path = str(reconciliation_xlsx)
    return report
