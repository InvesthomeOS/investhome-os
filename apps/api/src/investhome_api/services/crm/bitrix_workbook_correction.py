"""Apply the user-corrected Bitrix rescue workbook.

Does not invent phones, emails, units, or identities. Blank stays blank.
Wrong person on an agreement is repaired by relinking or creating a
source-backed review-required contact.
"""

from __future__ import annotations

from collections import defaultdict
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session, attributes

from investhome_api.models.crm_activity import CrmActivity, CrmActivityEntityType
from investhome_api.models.crm_agreement import CrmAgreement
from investhome_api.models.crm_contact import (
    CrmContact,
    CrmContactMergeHistory,
    CrmContactStatus,
    CrmContactType,
    CrmLifecycleStage,
    CrmRecordKind,
)
from investhome_api.models.sales import SalesOpportunity
from investhome_api.services.crm.bitrix_import import BitrixBundle, BitrixSourceRow, excel_number_as_text
from investhome_api.services.crm.bitrix_project_aliases import BitrixProjectGroup, fold_alias, resolve_project_group
from investhome_api.services.crm.bitrix_reit_membership import EXPECTED_AGREEMENT_COUNTS, EXPECTED_AGREEMENT_TOTAL
from investhome_api.services.crm.identity import (
    is_zero_masked_phone,
    normalize_full_name,
    normalize_valid_email,
    parse_phone,
    phone_is_better,
)

_METADATA_KEY = "bitrix_import"
_INTENTIONALLY_BLANK = {"", "—", "-", "n/a", "na", "none", "null", "yok"}


@dataclass
class WorkbookCorrectionReport:
    dry_run: bool
    detected_changes: list[dict[str, Any]] = field(default_factory=list)
    phones_corrected: list[dict[str, Any]] = field(default_factory=list)
    phones_left_blank: list[str] = field(default_factory=list)
    emails_names_corrected: list[dict[str, Any]] = field(default_factory=list)
    units_corrected: list[dict[str, Any]] = field(default_factory=list)
    links_repaired: list[dict[str, Any]] = field(default_factory=list)
    contacts_created: list[dict[str, Any]] = field(default_factory=list)
    lale_merge: dict[str, Any] | None = None
    selim_result: dict[str, Any] | None = None
    contacts_before: int = 0
    contacts_after: int = 0
    agreements_after: dict[str, int] = field(default_factory=dict)
    reit_count: int = 0
    notes: list[str] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)

    def format_console(self) -> str:
        payload = self.to_dict()
        lines = ["BITRIX WORKBOOK CORRECTION", ""]
        for key, value in payload.items():
            if isinstance(value, list) and key != "notes":
                lines.append(f"  {key}: {len(value)}")
            else:
                lines.append(f"  {key}: {value}")
        return "\n".join(lines) + "\n"


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _bitrix_meta(contact: CrmContact) -> dict[str, Any]:
    raw = contact.metadata_json if isinstance(contact.metadata_json, dict) else {}
    nested = raw.get(_METADATA_KEY)
    return dict(nested) if isinstance(nested, dict) else {}


def _set_bitrix_meta(contact: CrmContact, updates: dict[str, Any]) -> None:
    raw = dict(contact.metadata_json) if isinstance(contact.metadata_json, dict) else {}
    nested = dict(raw.get(_METADATA_KEY)) if isinstance(raw.get(_METADATA_KEY), dict) else {}
    nested.update(updates)
    raw[_METADATA_KEY] = nested
    contact.metadata_json = raw
    attributes.flag_modified(contact, "metadata_json")


def _cell_text(value: Any) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        text = excel_number_as_text(value)
        return text or None
    text = str(value).strip()
    return text or None


def _header_map(headers: list[Any]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    for index, header in enumerate(headers):
        key = fold_alias(str(header) if header is not None else "")
        if key and key not in mapped:
            mapped[key] = index
    return mapped


def _column(mapped: dict[str, int], aliases: set[str], row: list[Any]) -> str | None:
    for alias in aliases:
        index = mapped.get(alias)
        if index is None or index >= len(row):
            continue
        text = _cell_text(row[index])
        if text:
            return text
    return None


def _verified_phone(value: Any) -> str | None:
    text = _cell_text(value)
    if not text or fold_alias(text) in _INTENTIONALLY_BLANK:
        return None
    if is_zero_masked_phone(text):
        return None
    parsed = parse_phone(text)
    if parsed and parsed.e164 and not parsed.suspicious:
        return parsed.e164
    return None


def _normalize_unit(value: Any) -> str | None:
    text = _cell_text(value)
    if not text:
        return None
    if " / " in text:
        text = text.split(" / ", 1)[1].strip()
    folded = fold_alias(text)
    if folded in _INTENTIONALLY_BLANK:
        return None
    if folded.startswith("unit "):
        if any(token in folded for token in ("not applicable", "absent", "ambiguous", "genuinely")):
            return None
        rest = text.split(None, 1)[1] if " " in text else ""
        parts = [part.strip() for part in rest.replace("Unit", " ").replace("unit", " ").split(",") if part.strip()]
        if len(parts) == 1:
            return parts[0]
        if len(parts) == 2:
            return f"{parts[0]}-{parts[1]}"
        return "-".join(parts) if parts else None
    return text


def _sheet_rows(path: Path) -> dict[str, list[dict[str, Any]]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    out: dict[str, list[dict[str, Any]]] = {}
    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)
            try:
                headers = list(next(rows))
            except StopIteration:
                continue
            mapped = _header_map(headers)
            records: list[dict[str, Any]] = []
            for raw in rows:
                values = list(raw)
                if not any(_cell_text(item) for item in values):
                    continue
                records.append({"mapped": mapped, "values": values, "sheet": sheet.title})
            out[sheet.title] = records
    finally:
        workbook.close()
    return out


def _record_source_id(item: dict[str, Any]) -> str | None:
    mapped, values = item["mapped"], item["values"]
    return _column(
        mapped,
        {"bitrix external source id", "bitrix id", "source id", "external id", "external ids"},
        values,
    )


def _index_by_source_id(sheets: dict[str, list[dict[str, Any]]]) -> dict[str, list[dict[str, Any]]]:
    indexed: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for records in sheets.values():
        for item in records:
            source_id = _record_source_id(item)
            if source_id:
                indexed[source_id].append(item)
            extras = _column(item["mapped"], {"external ids"}, item["values"])
            if extras and extras != source_id:
                for part in str(extras).split(","):
                    key = part.strip()
                    if key:
                        indexed[key].append(item)
    return indexed


def _fold_source_id(value: str | None) -> str:
    text = (value or "").strip()
    if not text:
        return ""
    file_part, sep, bitrix_id = text.rpartition(":")
    if not sep:
        return fold_alias(text)
    return f"{fold_alias(file_part)}:{bitrix_id.strip()}"


def _lookup_folded(index: dict[str, Any], source_id: str | None) -> Any | None:
    if not source_id:
        return None
    if source_id in index:
        return index[source_id]
    needle = _fold_source_id(source_id)
    for key, value in index.items():
        if _fold_source_id(key) == needle:
            return value
    return None


def detect_workbook_changes(original: Path | None, current: Path) -> list[dict[str, Any]]:
    current_sheets = _sheet_rows(current)
    original_sheets = _sheet_rows(original) if original and original.exists() else {}
    current_ids = _index_by_source_id(current_sheets)
    original_ids = _index_by_source_id(original_sheets)
    changes: list[dict[str, Any]] = []
    for source_id in sorted(set(current_ids) | set(original_ids)):
        cur_items = current_ids.get(source_id, [])
        orig_items = original_ids.get(source_id, [])
        if not orig_items and cur_items:
            changes.append({"source_id": source_id, "kind": "added_row"})
            continue
        if orig_items and not cur_items:
            name = _column(orig_items[0]["mapped"], {"full name", "source full name", "display name"}, orig_items[0]["values"])
            changes.append({"source_id": source_id, "kind": "deleted_row", "name": name})
            continue
        cur = _best_correction_fields(cur_items)
        orig = _best_correction_fields(orig_items)
        delta = {key: {"from": orig.get(key), "to": cur.get(key)} for key in ("phone", "email", "name", "unit") if orig.get(key) != cur.get(key)}
        if delta:
            changes.append({"source_id": source_id, "kind": "cell_change", "fields": delta, "name": cur.get("name") or orig.get("name")})
    return changes


def _best_correction_fields(items: list[dict[str, Any]]) -> dict[str, str | None]:
    phone = email = name = unit = raw_phone = None
    for item in items:
        mapped, values = item["mapped"], item["values"]
        raw = _column(
            mapped,
            {"raw source phone", "current raw source phone", "corrected phone", "phone"},
            values,
        )
        if raw and not raw_phone:
            raw_phone = raw
        phone = phone or _verified_phone(raw)
        email = email or normalize_valid_email(_column(mapped, {"email", "e posta"}, values))
        name = name or _column(mapped, {"source full name", "full name"}, values)
        unit = unit or (
            _normalize_unit(_column(mapped, {"unit number"}, values))
            or _normalize_unit(_column(mapped, {"project unit"}, values))
            or _normalize_unit(_column(mapped, {"classification reason"}, values))
        )
    return {"phone": phone, "email": email, "name": name, "unit": unit, "raw_phone": raw_phone}


def _contacts_by_email(contacts: list[CrmContact]) -> dict[str, list[CrmContact]]:
    index: dict[str, list[CrmContact]] = defaultdict(list)
    for contact in contacts:
        for email in [contact.primary_email, *(contact.secondary_emails or [])]:
            key = normalize_valid_email(email)
            if key:
                index[key].append(contact)
    return index


def _contacts_by_external(contacts: list[CrmContact]) -> dict[str, CrmContact]:
    index: dict[str, CrmContact] = {}
    for contact in contacts:
        for value in _bitrix_meta(contact).get("external_ids") or []:
            key = str(value).strip()
            if key:
                index[key] = contact
    return index


def _append_external(contact: CrmContact, source_id: str, source_file: str | None) -> None:
    meta = _bitrix_meta(contact)
    ids = [str(item) for item in (meta.get("external_ids") or [])]
    files = [str(item) for item in (meta.get("source_files") or [])]
    roles = [str(item) for item in (meta.get("source_roles") or [])]
    if source_id not in ids:
        ids.append(source_id)
    if source_file and source_file not in files:
        files.append(source_file)
    if "agreements" not in roles:
        roles.append("agreements")
    _set_bitrix_meta(contact, {"external_ids": ids, "source_files": files, "source_roles": roles})


def _drop_external(contact: CrmContact, source_id: str) -> None:
    meta = _bitrix_meta(contact)
    ids = [str(item) for item in (meta.get("external_ids") or []) if str(item) != source_id]
    _set_bitrix_meta(contact, {"external_ids": ids})


def _create_source_contact(
    db: Session,
    *,
    name: str,
    source_id: str,
    source_file: str | None,
    email: str | None,
    phone: str | None,
) -> CrmContact:
    contact = CrmContact(
        contact_type=CrmContactType.BUYER,
        record_kind=CrmRecordKind.PERSON,
        display_name=name[:255],
        primary_email=email,
        primary_phone=phone,
        lifecycle_stage=CrmLifecycleStage.NEW,
        source="Bitrix",
        status=CrmContactStatus.ACTIVE,
        review_required=True,
        metadata_json={
            _METADATA_KEY: {
                "external_ids": [source_id],
                "source_files": [source_file] if source_file else [],
                "source_roles": ["agreements"],
                "created_from_workbook_correction": True,
                "created_at": _now(),
            }
        },
    )
    db.add(contact)
    db.flush()
    return contact


def _find_canonical(
    *,
    name: str | None,
    email: str | None,
    source_id: str | None,
    by_email: dict[str, list[CrmContact]],
    by_external: dict[str, CrmContact],
) -> CrmContact | None:
    if source_id and source_id in by_external:
        return by_external[source_id]
    if email:
        hits = by_email.get(email) or []
        named = [item for item in hits if name and normalize_full_name(item.display_name) == normalize_full_name(name)]
        if len(named) == 1:
            return named[0]
        active = [item for item in hits if item.status == CrmContactStatus.ACTIVE]
        if len(active) == 1:
            return active[0]
        if len(hits) == 1:
            return hits[0]
    return None


def _identity_wrong(source_name: str | None, source_email: str | None, contact: CrmContact) -> bool:
    src_name = normalize_full_name(source_name)
    dst_name = normalize_full_name(contact.display_name)
    src_email = normalize_valid_email(source_email)
    dst_email = normalize_valid_email(contact.primary_email)
    if not src_name or not dst_name or src_name == dst_name:
        return False
    if src_email and dst_email and src_email == dst_email:
        return False
    return True


def _merge_lale(db: Session, contacts: list[CrmContact], dry_run: bool) -> dict[str, Any] | None:
    lales = [
        item
        for item in contacts
        if normalize_full_name(item.display_name) == normalize_full_name("Lale Şenyol")
        and normalize_valid_email(item.primary_email) == "lsenyol@gmail.com"
        and item.status == CrmContactStatus.ACTIVE
    ]
    if len(lales) < 2:
        return {"status": "already_single", "count": len(lales)}
    by_id = {item.id: item for item in lales}
    agreements = list(db.query(CrmAgreement).filter(CrmAgreement.contact_id.in_(list(by_id))).all())
    survivor = lales[0]
    for contact in lales:
        ids = " ".join(_bitrix_meta(contact).get("external_ids") or [])
        if "1812" in ids:
            survivor = contact
            break
    merged = [item for item in lales if item.id != survivor.id]
    payload = {
        "survivor_id": str(survivor.id),
        "merged_ids": [str(item.id) for item in merged],
        "agreements_moved": 0,
    }
    if dry_run:
        return payload
    for other in merged:
        for agreement in db.query(CrmAgreement).filter(CrmAgreement.contact_id == other.id):
            agreement.contact_id = survivor.id
            payload["agreements_moved"] += 1
        for activity in db.query(CrmActivity).filter(
            CrmActivity.entity_type == CrmActivityEntityType.CONTACT,
            CrmActivity.entity_id == other.id,
        ):
            activity.entity_id = survivor.id
        for opportunity in db.query(SalesOpportunity).filter(SalesOpportunity.crm_contact_id == other.id):
            opportunity.crm_contact_id = survivor.id
        other_meta = _bitrix_meta(other)
        survivor_meta = _bitrix_meta(survivor)
        ids = list(dict.fromkeys([*(survivor_meta.get("external_ids") or []), *(other_meta.get("external_ids") or [])]))
        files = list(dict.fromkeys([*(survivor_meta.get("source_files") or []), *(other_meta.get("source_files") or [])]))
        roles = list(dict.fromkeys([*(survivor_meta.get("source_roles") or []), *(other_meta.get("source_roles") or [])]))
        _set_bitrix_meta(survivor, {"external_ids": ids, "source_files": files, "source_roles": roles, "lale_merged": True})
        if other.primary_phone and other.primary_phone != survivor.primary_phone:
            if not survivor.primary_phone:
                survivor.primary_phone = other.primary_phone
            else:
                extras = list(survivor.secondary_phones or [])
                if other.primary_phone not in extras:
                    extras.append(other.primary_phone)
                    survivor.secondary_phones = extras
        if other.primary_email and other.primary_email != survivor.primary_email:
            if not survivor.primary_email:
                survivor.primary_email = other.primary_email
            else:
                extras = list(survivor.secondary_emails or [])
                if other.primary_email not in extras:
                    extras.append(other.primary_email)
                    survivor.secondary_emails = extras
        if other.notes and (other.notes not in (survivor.notes or "")):
            survivor.notes = f"{survivor.notes}\n{other.notes}".strip() if survivor.notes else other.notes
        db.add(
            CrmContactMergeHistory(
                survivor_contact_id=survivor.id,
                merged_contact_id=other.id,
                merged_snapshot={"display_name": other.display_name, "primary_email": other.primary_email},
                merged_by_user_id=None,
            )
        )
        other.status = CrmContactStatus.ARCHIVED
        other.archived_at = datetime.now(timezone.utc)
        other.review_required = False
    del agreements
    return payload


def run_bitrix_workbook_correction(
    db: Session,
    bundle: BitrixBundle,
    *,
    rescue_xlsx: Path,
    original_xlsx: Path | None = None,
    dry_run: bool = True,
) -> WorkbookCorrectionReport:
    contacts = list(db.query(CrmContact).all())
    agreements = list(db.query(CrmAgreement).all())
    report = WorkbookCorrectionReport(dry_run=dry_run, contacts_before=len(contacts))
    report.detected_changes = detect_workbook_changes(original_xlsx, rescue_xlsx)
    current_sheets = _sheet_rows(rescue_xlsx)
    fields_by_id: dict[str, dict[str, str | None]] = {}
    for source_id, items in _index_by_source_id(current_sheets).items():
        fields_by_id[source_id] = _best_correction_fields(items)

    by_email = _contacts_by_email(contacts)
    by_external = _contacts_by_external(contacts)
    by_source_agreement = {
        row.source_external_id: row for row in agreements if row.source_external_id
    }
    source_rows = {
        f"{row.source_file}:{row.bitrix_id}": row
        for row in bundle.rows
        if row.role == "agreements" and row.bitrix_id
    }

    # Repair wrong person links first so phones land on the correct contact.
    for source_id, source_row in source_rows.items():
        agreement = _lookup_folded(by_source_agreement, source_id)
        if agreement is None:
            continue
        contact = next((item for item in contacts if item.id == agreement.contact_id), None)
        if contact is None:
            continue
        workbook = _lookup_folded(fields_by_id, source_id) or {}
        source_name = source_row.full_name
        source_email = normalize_valid_email(source_row.email)
        if not _identity_wrong(source_name, source_email, contact):
            continue
        target = _find_canonical(
            name=source_name,
            email=source_email,
            source_id=None,
            by_email=by_email,
            by_external=by_external,
        )
        if target is not None and normalize_full_name(target.display_name) != normalize_full_name(source_name):
            if source_email and normalize_valid_email(target.primary_email) != source_email:
                target = None
        created = False
        if target is None:
            phone = workbook.get("phone")
            if dry_run:
                target_id = None
            else:
                target = _create_source_contact(
                    db,
                    name=source_name or source_id,
                    source_id=source_id,
                    source_file=source_row.source_file,
                    email=source_email,
                    phone=phone,
                )
                contacts.append(target)
                by_external[source_id] = target
                if source_email:
                    by_email[source_email].append(target)
                target_id = target.id
            created = True
            report.contacts_created.append(
                {
                    "source_id": source_id,
                    "name": source_name,
                    "email": source_email,
                    "phone": phone,
                    "contact_id": str(target_id) if target_id else None,
                }
            )
        if not dry_run and target is not None and target.id != agreement.contact_id:
            _drop_external(contact, source_id)
            _append_external(target, source_id, source_row.source_file)
            agreement.contact_id = target.id
        report.links_repaired.append(
            {
                "source_id": source_id,
                "from": contact.display_name,
                "to": source_name,
                "created": created,
            }
        )

    if not dry_run:
        db.flush()
        contacts = list(db.query(CrmContact).all())
        agreements = list(db.query(CrmAgreement).all())
        by_email = _contacts_by_email(contacts)
        by_external = _contacts_by_external(contacts)
        by_source_agreement = {row.source_external_id: row for row in agreements if row.source_external_id}

    for source_id, fields in fields_by_id.items():
        agreement = _lookup_folded(by_source_agreement, source_id)
        contact = None
        if agreement is not None:
            contact = next((item for item in contacts if item.id == agreement.contact_id), None)
        if contact is None:
            contact = _lookup_folded(by_external, source_id)
        if contact is None:
            continue
        stale = _lookup_folded(by_external, source_id)
        if not dry_run and stale is not None and stale.id != contact.id:
            _drop_external(stale, source_id)
            source_row_for_meta = _lookup_folded(source_rows, source_id)
            _append_external(contact, source_id, source_row_for_meta.source_file if source_row_for_meta else None)
            by_external[source_id] = contact
        source_row = _lookup_folded(source_rows, source_id)
        group = resolve_project_group(source_row.source_file) if source_row is not None else None
        phone = fields.get("phone")
        if phone and (contact.primary_phone or "") != phone:
            if not dry_run:
                contact.primary_phone = phone
            report.phones_corrected.append(
                {"source_id": source_id, "name": contact.display_name, "phone": phone}
            )
        elif not phone and agreement is not None and (
            is_zero_masked_phone(fields.get("raw_phone")) or is_zero_masked_phone(contact.primary_phone)
        ):
            if contact.display_name not in report.phones_left_blank:
                report.phones_left_blank.append(contact.display_name)
            if not dry_run and is_zero_masked_phone(contact.primary_phone):
                contact.primary_phone = None
        email = fields.get("email")
        source_email = normalize_valid_email(source_row.email) if source_row is not None else None
        if source_row is not None and email and source_email and email != source_email:
            email = None
        if email and not contact.primary_email:
            if not dry_run:
                contact.primary_email = email
            report.emails_names_corrected.append(
                {"source_id": source_id, "field": "email", "value": email}
            )
        name = fields.get("name")
        if source_row is not None:
            name = source_row.full_name or name
        if (
            name
            and normalize_full_name(name) != normalize_full_name(contact.display_name)
            and not _identity_wrong(name, source_email, contact)
        ):
            if not dry_run:
                contact.display_name = name[:255]
            report.emails_names_corrected.append(
                {"source_id": source_id, "field": "name", "value": name}
            )
        unit = fields.get("unit")
        if agreement is not None and group != BitrixProjectGroup.REIT and unit and (agreement.unit_number or "") != unit:
            previous = agreement.unit_number
            if not dry_run:
                agreement.unit_number = unit[:80]
            report.units_corrected.append(
                {"source_id": source_id, "from": previous, "to": unit}
            )
        if agreement is not None and group == BitrixProjectGroup.REIT and agreement.unit_number:
            if not dry_run:
                agreement.unit_number = None

    # Sheet 3 junk/forensic phone+email fills/clears keyed by external id.
    for item in current_sheets.get("3 Contact delta +153", []):
        source_id = _record_source_id(item)
        if not source_id:
            continue
        contact = by_external.get(source_id)
        if contact is None:
            continue
        mapped, values = item["mapped"], item["values"]
        new_phone = _column(mapped, {"phone"}, values)
        new_email = normalize_valid_email(_column(mapped, {"email"}, values))
        verified = _verified_phone(new_phone)
        if verified and phone_is_better(verified, contact.primary_phone):
            if not dry_run:
                contact.primary_phone = verified
            report.phones_corrected.append(
                {"source_id": source_id, "name": contact.display_name, "phone": verified, "sheet": "3"}
            )
        elif new_phone is None and contact.primary_phone and is_zero_masked_phone(contact.primary_phone):
            if not dry_run:
                contact.primary_phone = None
            if source_id not in report.phones_left_blank:
                report.phones_left_blank.append(source_id)
        if new_email and not contact.primary_email:
            if not dry_run:
                contact.primary_email = new_email
            report.emails_names_corrected.append(
                {"source_id": source_id, "field": "email", "value": new_email, "sheet": "3"}
            )

    lale = _merge_lale(db, contacts, dry_run=dry_run)
    report.lale_merge = lale

    selim_agreements = [
        row for row in (db.query(CrmAgreement).all() if not dry_run else agreements)
        if row.source_external_id == "Anlaşmalar Reıt.xls:238"
    ]
    if selim_agreements:
        contact = db.get(CrmContact, selim_agreements[0].contact_id)
        report.selim_result = {
            "contact_id": str(selim_agreements[0].contact_id),
            "contact_name": contact.display_name if contact else None,
            "email": contact.primary_email if contact else None,
            "source_id": "Anlaşmalar Reıt.xls:238",
        }

    if not dry_run:
        db.flush()
        contacts = list(db.query(CrmContact).all())
        agreements = list(db.query(CrmAgreement).all())
    report.contacts_after = len(contacts)
    counts: dict[str, int] = {}
    for row in agreements:
        counts[row.project_group] = counts.get(row.project_group, 0) + 1
    report.agreements_after = dict(sorted(counts.items()))
    report.reit_count = counts.get(BitrixProjectGroup.REIT.value, 0)
    if sum(counts.values()) != EXPECTED_AGREEMENT_TOTAL:
        report.notes.append(f"agreement_total:{sum(counts.values())}")
    for group, expected in EXPECTED_AGREEMENT_COUNTS.items():
        actual = counts.get(group, 0)
        if actual != expected:
            report.notes.append(f"count_mismatch:{group}:{actual}:{expected}")
    return report
